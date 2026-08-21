"""One-time, idempotent fix: the "Đại diện Bên A/B/C" signature block in
the transfer template was three absolutely-positioned floating text boxes
(`wp:anchor`), not a table -- fragile to edit by hand and easy to break by
just adding a line of text above it (the boxes don't reflow). Replaces them
with a real 3-column table in the same spot, carrying over the exact same
placeholder tokens (fixing a duplicated-placeholder bug in Bên A's box
along the way), and adds the provider's stamp+signature image to Bên B --
the one party this app can supply a real signature image for.

Run from the repository root with the desktop virtual environment.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "desktop_app/backend/documents/transfer/00_MAU_BIEN_BAN_CHUYEN_CHU_QUYEN.docx"
PROVIDER_SIGNATURE = ROOT / "desktop_app/data/source/signatures/vietnamobile_provider_stamp_signature.png"

COLUMNS = [
    ("Đại diện Bên A", "{{ customer_representative_signature_name }}", False),
    ("Đại diện Bên B", "{{ provider_representative }}", True),
    ("Đại diện Bên C", "{{ new_owner_signature_name }}", False),
]
HINT = "(Ký, đóng dấu, ghi họ tên)"


def _set_cell_margins(cell, top=80, bottom=80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("bottom", bottom)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def _write(paragraph, text: str, *, size=10, bold=False, italic=False) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def _signature_cell(cell, heading: str, content: str, *, provider: bool) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    _set_cell_margins(cell)
    _write(cell.paragraphs[0], heading, bold=True)
    hint = cell.add_paragraph()
    _write(hint, HINT, size=8.5, italic=True)
    body = cell.add_paragraph()
    body.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if provider and PROVIDER_SIGNATURE.is_file():
        body.add_run().add_picture(str(PROVIDER_SIGNATURE), width=Mm(45))
        body.add_run().add_break()
    run = body.add_run(content)
    run.font.size = Pt(9.5)
    run.font.bold = True


def build(template: Path = TEMPLATE) -> int:
    document = Document(str(template))
    host = next(
        (p for p in document.paragraphs if p._p.findall(".//" + qn("wp:anchor"))
         and len(p._p.findall(".//" + qn("wp:anchor"))) == 3),
        None,
    )
    if host is None:
        return 0

    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, (heading, content, provider) in zip(table.rows[0].cells, COLUMNS):
        _signature_cell(cell, heading, content, provider=provider)

    host._p.addnext(table._tbl)
    host._p.getparent().remove(host._p)
    document.save(template)
    return 1


if __name__ == "__main__":
    print(f"replaced signature block: {bool(build())}")
