"""One-time, idempotent fix: Aftersale's "NGƯỜI YÊU CẦU / CHỦ THUÊ BAO MỚI /
GIAO DỊCH VIÊN" signature block was a real Word 3-column SECTION (`w:cols`)
with the three columns' content simply flowing linearly through one
sequence of paragraphs -- the same class of fragility (implicit, layout-
engine-dependent positioning instead of an explicit table) that the other
4 templates' signature blocks were already rebuilt away from earlier this
session. Converts it to a real 3-row x 3-col table: heading+hint / a
fixed-height signature row (blank for requester/new-owner, the clerk's
{{operator_signature_marker}} for the third column) / the printed name --
same pattern as transfer's table, so the three printed names land on the
same Y position regardless of which column carries a signature image.

The final section (the ONLY thing using 3 columns) is reset to a single
column, since that setting existed purely to fake this layout.

Run from the repository root with the desktop virtual environment.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "desktop_app/backend/documents/aftersale/00_MAU_CAM_KET_SAU_BAN_HANG.docx"
SIGNATURE_ROW_HEIGHT = Mm(32)

COLUMNS = [
    ("NGƯỜI YÊU CẦU", "{{ aftersale_requester_signature_name }}", False),
    ("CHỦ THUÊ BAO MỚI", "{{ aftersale_new_owner_signature_name }}", False),
    ("GIAO DỊCH VIÊN", "{{ aftersale_clerk_signature_name }}", True),
]
HINT = "(Ký và ghi rõ họ tên)"
MARKER = "{{operator_signature_marker}}"


def _set_cell_width(cell, width_mm: float) -> None:
    # Word's default column autofit left this table badly lopsided in
    # practice (one column visible, the other two squeezed to near-zero or
    # pushed past the page edge) -- explicit fixed widths on every cell,
    # not just the header row, is what actually holds across Word and
    # LibreOffice's rendering.
    cell.width = Mm(width_mm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(round(width_mm * 56.6929)))
    tc_w.set(qn("w:type"), "dxa")


def _set_cell_margins(cell, top=60, bottom=60) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("bottom", bottom)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def _write(paragraph, text: str, *, size=10, bold=False, italic=False) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def build(template: Path = TEMPLATE) -> bool:
    document = Document(str(template))
    heading_index = next(
        (i for i, p in enumerate(document.paragraphs) if p.text.strip() == "NGƯỜI YÊU CẦU"),
        None,
    )
    if heading_index is None:
        return False
    paragraphs = document.paragraphs
    host = paragraphs[heading_index]._p

    # The final section is entirely this signature block -- drop its w:cols
    # entirely now that a real table replaces the 3-column flow (absence
    # means Word's true single-column default). Just setting num="1" isn't
    # enough: the element still carries 3 leftover <w:col> width/spacing
    # children from the original definition, and LibreOffice's renderer
    # follows those over num, squeezing the table into one of the three
    # narrow text regions instead of the full page width.
    final_section_properties = document.sections[-1]._sectPr
    columns = final_section_properties.find(qn("w:cols"))
    if columns is not None:
        final_section_properties.remove(columns)

    table = document.add_table(rows=3, cols=len(COLUMNS))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    table.rows[1].height = SIGNATURE_ROW_HEIGHT
    table.rows[1].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY

    usable_mm = (
        document.sections[-1].page_width
        - document.sections[-1].left_margin
        - document.sections[-1].right_margin
    ) / 36000
    column_width_mm = usable_mm / len(COLUMNS)

    for col_index, (heading, name_token, is_clerk) in enumerate(COLUMNS):
        for row_index in range(3):
            _set_cell_width(table.cell(row_index, col_index), column_width_mm)

        heading_cell = table.cell(0, col_index)
        heading_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_margins(heading_cell)
        _write(heading_cell.paragraphs[0], heading, bold=True)
        hint_p = heading_cell.add_paragraph()
        _write(hint_p, HINT, size=8.5, italic=True)

        signature_cell = table.cell(1, col_index)
        signature_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_margins(signature_cell, top=20, bottom=20)
        if is_clerk:
            _write(signature_cell.paragraphs[0], MARKER, size=9, bold=True)

        name_cell = table.cell(2, col_index)
        name_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        _set_cell_margins(name_cell, top=10, bottom=60)
        _write(name_cell.paragraphs[0], name_token, size=9.5, bold=True)

    # Remove the old heading paragraph and everything after it (the rest of
    # the linear 3-column flow) -- the table takes its place. The body's own
    # trailing sectPr (page setup for the final section) is a separate XML
    # node, not one of these paragraphs, so it's untouched.
    trailing = host.getnext()
    host.addprevious(table._tbl)
    to_remove = [host]
    while trailing is not None and trailing.tag == qn("w:p"):
        next_node = trailing.getnext()
        to_remove.append(trailing)
        trailing = next_node
    for node in to_remove:
        node.getparent().remove(node)

    document.save(template)
    return True


if __name__ == "__main__":
    print(f"rebuilt aftersale signature table: {build()}")
