"""Repair collateral damage from remove_baked_dots.py.

That script's regex (`\\.{5}\\s?|\\s?\\.{5}`) removed 5-dot groups anywhere
in a run's text, on the assumption every "....." in these templates was
bake_field_highlighting.py's placeholder decoration. Two template files
turned out to also contain STATIC, placeholder-less dotted blank lines
(a physical "write the change by hand" convention, unrelated to any
{{ }} token) whose dot counts weren't multiples of 5 -- so those runs got
partially eaten instead of matched cleanly, and in a couple of spots a
run's adjacent newline got consumed along with its dots (the pattern's
`\\s?` matches ANY single whitespace character, not just a space).

Found by comparing each damaged line's label text against a "bound"
sibling paragraph carrying the same label with a real placeholder still
attached, then re-deriving a blank length from renderer.py's own
`_empty_field_placeholder` length table (name=28, address=36, id/phone/
number=16, date=12, email=24, else=16) -- the SAME convention this
codebase already uses everywhere else for "how long should an empty
field's blank be," so the restored lengths are principled, not guessed.
The exact ORIGINAL dot counts are not recoverable (no backup existed);
this restores each blank to that established convention's length rather
than reproducing byte-for-byte original content.

sim_change_form/00_MAU_PHIEU_THAY_DOI_DICH_VU_TRA_TRUOC.docx: paragraphs
22/24/26/28/30/32/34 are the "Thông tin khách hàng thay đổi (Nếu có)"
column -- a static, intentionally UNBOUND mirror of the adjacent "đã
cung cấp" column's real placeholders (customer hand-writes a correction
there if anything changed) -- rebuilt each from scratch, textually
matching its bound sibling's label wording and line breaks exactly.

aftersale/00_MAU_CAM_KET_SAU_BAN_HANG.docx: paragraph 5's "Shop Xã Đàn"
(a real, complete value, not a blank -- trailing dots served no purpose)
just has its stray ".." remnant dropped. Paragraph 12's 2 trailing-dots
markers (after aftersale_customer_id_number / aftersale_customer_issue_date,
both label-preceded so no leading dots, per bake_field_highlighting.py's
own is_label rule) are restored to the standard " ..... " marker.
Paragraphs 6/7 ("Địa chỉ:"/"Điện thoại:") show no leftover dot fragment
at all, so whether they ever had a dotted blank can't be determined from
current state -- left untouched rather than fabricated.

Run from the repository root with the desktop virtual environment:
    python desktop_app/scripts/repair_dot_removal_damage.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from desktop_app.backend.documents.renderer import _iter_paragraphs

ROOT = Path(__file__).resolve().parents[2]
SIM_CHANGE_PATH = ROOT / "desktop_app/backend/documents/sim_change_form/00_MAU_PHIEU_THAY_DOI_DICH_VU_TRA_TRUOC.docx"
AFTERSALE_PATH = ROOT / "desktop_app/backend/documents/aftersale/00_MAU_CAM_KET_SAU_BAN_HANG.docx"

NAME = "." * 28
ADDRESS = "." * 36
ID_NUMBER = "." * 16
DATE = "." * 12
EMAIL = "." * 24
OTHER = "." * 16

# (paragraph text to find (by exact old, damaged content), replacement text)
SIM_CHANGE_FIXES = [
    (
        "Tên cơ quan, tổ chức hoặc cá nhân (viết in hoa):....",
        f"Tên cơ quan, tổ chức hoặc cá nhân (viết in hoa):\n{NAME}",
    ),
    (
        "Địa chỉ trụ sở chính/Địa chỉ theo CCCD/Căn cước/Hộ chiếu:....",
        f"Địa chỉ trụ sở chính/Địa chỉ theo CCCD/Căn cước/Hộ chiếu:\n{ADDRESS}",
    ),
    (
        "- Số QĐTL/GCNĐKKD&ĐKĐT/GPĐT/GCNĐKDN:- Nơi cấp/Đơn vị cấp:- Ngày cấp:..",
        f"- Số QĐTL/GCNĐKKD&ĐKĐT/GPĐT/GCNĐKDN:\n{ID_NUMBER}"
        f"\n- Nơi cấp/Đơn vị cấp: {NAME}\n- Ngày cấp: {DATE}",
    ),
    (
        "- Người đại diện/ủy quyền:.\n- Số ĐDCN/ĐDĐT/Hộ chiếu:....\n"
        "- Ngày tháng năm sinh:- Giới tính:.\n- Ngày cấp:.\n"
        "- Nơi cấp/Đơn vị cấp:...\n"
        "- Địa chỉ theo giấy tờ dùng để đăng ký thông tin thuê bao:"
        "- Quốc tịch: ☐ Việt Nam    ☐ Nước ngoài",
        f"- Người đại diện/ủy quyền: {NAME}\n- Số ĐDCN/ĐDĐT/Hộ chiếu: {ID_NUMBER}"
        f"\n- Ngày tháng năm sinh: {DATE}\n- Giới tính: {OTHER}\n- Ngày cấp: {DATE}"
        f"\n- Nơi cấp/Đơn vị cấp: {NAME}"
        f"\n- Địa chỉ theo giấy tờ dùng để đăng ký thông tin thuê bao: {ADDRESS}"
        "\n- Quốc tịch: ☐ Việt Nam    ☐ Nước ngoài",
    ),
    (
        "- Nơi gửi thông báo cước và thanh toán:",
        f"- Nơi gửi thông báo cước và thanh toán: {ADDRESS}",
    ),
    (
        "- Số điện thoại liên hệ:",
        f"- Số điện thoại liên hệ: {ID_NUMBER}",
    ),
    (
        "- Email:.",
        f"- Email: {EMAIL}",
    ),
]

AFTERSALE_FIXES = [
    ("\t\tCửa hàng: Shop Xã Đàn..", "\t\tCửa hàng: Shop Xã Đàn"),
]

# Paragraph 12 mixes real, already-baked placeholder runs (bold/size set
# by bake_field_highlighting.py) with plain connector-text runs between
# them -- rewriting the whole paragraph via _rewrite_paragraph would wipe
# the placeholder runs' baked formatting, so these 2 connector runs are
# patched in place (run.text edit only, formatting untouched) instead.
AFTERSALE_RUN_FIXES = [
    (".   Ngày cấp: ", " ..... Ngày cấp: "),
    (".    Nơi cấp: ", " ..... Nơi cấp: "),
]


def _rewrite_paragraph(paragraph, new_text: str) -> None:
    for run in list(paragraph.runs):
        run._r.getparent().remove(run._r)
    lines = new_text.split("\n")
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        if line:
            paragraph.add_run(line)


def _apply_fixes(path: Path, fixes: list[tuple[str, str]]) -> int:
    document = Document(str(path))
    remaining = dict(fixes)
    applied = 0
    for paragraph in _iter_paragraphs(document):
        if paragraph.text in remaining:
            _rewrite_paragraph(paragraph, remaining.pop(paragraph.text))
            applied += 1
    document.save(str(path))
    if remaining:
        print(f"  NOT FOUND (already fixed or text changed): {list(remaining)}")
    return applied


def _apply_run_fixes(path: Path, fixes: list[tuple[str, str]]) -> int:
    document = Document(str(path))
    remaining = dict(fixes)
    applied = 0
    for paragraph in _iter_paragraphs(document):
        for run in paragraph.runs:
            if run.text in remaining:
                run.text = remaining.pop(run.text)
                applied += 1
    document.save(str(path))
    if remaining:
        print(f"  NOT FOUND (already fixed or text changed): {list(remaining)}")
    return applied


if __name__ == "__main__":
    count = _apply_fixes(SIM_CHANGE_PATH, SIM_CHANGE_FIXES)
    print(f"{SIM_CHANGE_PATH.name}: repaired {count} paragraph(s)")
    count = _apply_fixes(AFTERSALE_PATH, AFTERSALE_FIXES)
    count += _apply_run_fixes(AFTERSALE_PATH, AFTERSALE_RUN_FIXES)
    print(f"{AFTERSALE_PATH.name}: repaired {count} paragraph(s)/run(s)")
