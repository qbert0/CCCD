"""One-time, idempotent fix: Beautiful Number and Prepaid Contract each
have a "Bên B" (Vietnamobile) signature slot that only ever showed the
literal text "VÕ DUY NHẬT" (never data-driven), and neither carried the
provider's actual stamp+signature image the way sim_change_form's rebuilt
template now does. Replaces the literal name with the real
{{ provider_representative }} placeholder and inserts the same stamp
image used elsewhere in the app.

Run from the repository root with the desktop virtual environment.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Mm, Pt

ROOT = Path(__file__).resolve().parents[2]
PROVIDER_SIGNATURE = ROOT / "desktop_app/data/source/signatures/vietnamobile_provider_stamp_signature.png"
HARDCODED_NAME = "VÕ DUY NHẬT"
PLACEHOLDER = "{{ provider_representative }}"


def _insert_signature(paragraph: Paragraph, *, width_mm: float) -> None:
    for run in paragraph.runs:
        run.text = ""
    picture_run = paragraph.add_run()
    picture_run.add_picture(str(PROVIDER_SIGNATURE), width=Mm(width_mm))
    paragraph.add_run().add_break()
    name_run = paragraph.add_run(PLACEHOLDER)
    name_run.font.bold = True
    name_run.font.size = Pt(9.5)


def _textbox_paragraphs(document: Document):
    for paragraph in document.paragraphs:
        for txbx_content in paragraph._p.findall(".//" + qn("w:txbxContent")):
            for p_element in txbx_content.findall(qn("w:p")):
                yield Paragraph(p_element, paragraph._parent)


def fix_beautiful_number(path: Path) -> int:
    document = Document(str(path))
    fixed = 0
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph.text.strip() == HARDCODED_NAME:
                        _insert_signature(paragraph, width_mm=42)
                        fixed += 1
    if fixed:
        document.save(path)
    return fixed


def fix_prepaid_contract(path: Path) -> int:
    document = Document(str(path))
    fixed = 0
    for paragraph in _textbox_paragraphs(document):
        if paragraph.text.strip() == HARDCODED_NAME:
            _insert_signature(paragraph, width_mm=42)
            fixed += 1
    if fixed:
        document.save(path)
    return fixed


def main() -> None:
    beautiful_number = ROOT / "desktop_app/backend/documents/beautiful_number/00_MAU_PHU_LUC_CAM_KET_SO_DEP_editable.docx"
    prepaid_contract = ROOT / "desktop_app/backend/documents/prepaid_contract/00_MAU_HOP_DONG_TRA_TRUOC.docx"
    print(f"{beautiful_number.name}: fixed {fix_beautiful_number(beautiful_number)} occurrence(s)")
    print(f"{prepaid_contract.name}: fixed {fix_prepaid_contract(prepaid_contract)} occurrence(s)")


if __name__ == "__main__":
    main()
