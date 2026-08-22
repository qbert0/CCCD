"""One-time, idempotent fix: transfer, prepaid_contract and beautiful_number
each had the provider's stamp+signature baked into the template as a
static picture (added earlier this session). Splits the paragraph holding
that picture into two: a {{provider_signature_marker}} paragraph (filled
in at generation time by renderer.py::_fill_provider_signature_marker from
the representative profile's own configurable signature_path) and the
existing {{ provider_representative }} name paragraph, unchanged. Lets the
shop swap in its own actual signature image via Thiết lập mặc định instead
of always printing whichever picture happened to be embedded here.

Run from the repository root with the desktop virtual environment.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[2]
TARGETS = [
    ROOT / "desktop_app/backend/documents/transfer/00_MAU_BIEN_BAN_CHUYEN_CHU_QUYEN.docx",
    ROOT / "desktop_app/backend/documents/prepaid_contract/00_MAU_HOP_DONG_TRA_TRUOC.docx",
    ROOT / "desktop_app/backend/documents/beautiful_number/00_MAU_PHU_LUC_CAM_KET_SO_DEP_editable.docx",
]
MARKER = "{{provider_signature_marker}}"


def _paragraph_has_picture(paragraph) -> bool:
    return bool(paragraph._p.findall(".//" + qn("w:drawing")))


def fix(path: Path) -> int:
    document = Document(str(path))
    fixed = 0
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in list(cell.paragraphs):
                    if not _paragraph_has_picture(paragraph):
                        continue
                    # Remove the picture run(s) and any trailing <w:br/> run
                    # left over from the old "picture, break, name" layout,
                    # keeping only the name run(s) in this paragraph.
                    for run in list(paragraph.runs):
                        if run._r.findall(".//" + qn("w:drawing")) or run._r.findall(".//" + qn("w:br")):
                            run._r.getparent().remove(run._r)
                    marker_p_element = OxmlElement("w:p")
                    paragraph._p.addprevious(marker_p_element)
                    from docx.text.paragraph import Paragraph
                    marker_p = Paragraph(marker_p_element, paragraph._parent)
                    marker_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    marker_run = marker_p.add_run(MARKER)
                    marker_run.font.size = Pt(9)
                    marker_run.font.bold = True
                    fixed += 1
    if fixed:
        document.save(path)
    return fixed


def main() -> None:
    for path in TARGETS:
        print(f"{path.name}: split {fix(path)} signature paragraph(s)")


if __name__ == "__main__":
    main()
