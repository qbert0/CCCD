"""One-time, idempotent migration of editable DOCX templates to placeholders.

Run from the repository root with the desktop virtual environment.  The generated
templates remain normal Word files: users can edit their layout and formatting as
long as the ``{{ placeholder_name }}`` tokens are retained.
"""

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[2]


def set_text(paragraph: Paragraph, text: str) -> None:
    """Change visible text while retaining paragraph and first-run formatting."""
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def set_text_after_column_break(paragraph: Paragraph, text: str) -> None:
    """Keep the section's column break before replacing its visible heading."""
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].add_break(WD_BREAK.COLUMN)
    else:
        paragraph.add_run().add_break(WD_BREAK.COLUMN)
    paragraph.add_run(text)


def replace_tokens(paragraph: Paragraph, replacements: dict[str, str]) -> None:
    """Rename placeholders even when Word has split a token across runs."""
    runs = paragraph.runs
    joined = "".join(run.text for run in runs)
    for old, new in replacements.items():
        while old in joined:
            start = joined.index(old)
            end = start + len(old)
            positions: list[tuple[int, int]] = []
            cursor = 0
            for run in runs:
                positions.append((cursor, cursor + len(run.text)))
                cursor += len(run.text)
            start_run = next(index for index, (_, stop) in enumerate(positions) if start < stop)
            end_run = next(index for index, (begin, stop) in enumerate(positions) if begin < end <= stop)
            start_offset = start - positions[start_run][0]
            end_offset = end - positions[end_run][0]
            prefix = runs[start_run].text[:start_offset]
            suffix = runs[end_run].text[end_offset:]
            if start_run == end_run:
                runs[start_run].text = prefix + new + suffix
            else:
                runs[start_run].text = prefix + new
                for index in range(start_run + 1, end_run):
                    runs[index].text = ""
                runs[end_run].text = suffix
            joined = "".join(run.text for run in runs)


def migrate_transfer() -> None:
    source = ROOT / (
        "desktop_app/backend/documents/transfer/"
        "Biên bản chuyển chủ quyền 2025.docx"
    )
    path = ROOT / "desktop_app/backend/documents/transfer/00_MAU_BIEN_BAN_CHUYEN_CHU_QUYEN.docx"
    document = Document(source)
    paragraphs = document.paragraphs
    common = {
        "{{pay_method}}": "{{ payment_method }}",
        "{{hop_dong_so}}": "{{ source_contract_number }}",
        "{{date_1}}": "{{ source_contract_day }}",
        "{{month_1}}": "{{ source_contract_month }}",
        "{{year_1}}": "{{ source_contract_year }}",
        "{{date_2}}": "{{ registration_form_day }}",
        "{{month_2}}": "{{ registration_form_month }}",
        "{{year_2}}": "{{ registration_form_year }}",
        "{{number}}": "{{ subscriber_number }}",
        "{{time}}": "{{ transfer_time }}",
    }
    for paragraph in paragraphs:
        replace_tokens(paragraph, common)
    replace_tokens(
        paragraphs[7],
        {
            "{{date_3}}": "{{ document_day }}",
            "{{month_3}}": "{{ document_month }}",
            "{{year_3}}": "{{ document_year }}",
        },
    )
    replace_tokens(
        paragraphs[12],
        {
            "{{date_3}}": "{{ transfer_effective_day }}",
            "{{month_3}}": "{{ transfer_effective_month }}",
            "{{year_3}}": "{{ transfer_effective_year }}",
        },
    )
    # The editable file used wide character spacing for the blank time field.
    # A real value such as 10:30 must remain normally spaced when printed.
    for run in paragraphs[12].runs:
        if "{{ transfer_time }}" not in run.text or run._r.rPr is None:
            continue
        spacing = run._r.rPr.find(qn("w:spacing"))
        if spacing is not None:
            run._r.rPr.remove(spacing)

    document.save(path)


def migrate_aftersale() -> None:
    path = ROOT / "desktop_app/backend/documents/aftersale/00_MAU_CAM_KET_SAU_BAN_HANG.docx"
    # The canonical file is intentionally the source: it contains the user's
    # improved visual design.  The original document is only a layout reference
    # and must never overwrite those edits.
    document = Document(path)
    paragraphs = document.paragraphs
    values = {
        3: "Ngày {{ aftersale_day }} tháng {{ aftersale_month }} năm {{ aftersale_year }}",
        5: "Cửa hàng: {{ aftersale_shop_name }}",
        6: "Địa chỉ: {{ aftersale_shop_address }}",
        7: "Điện thoại: {{ aftersale_shop_phone }}",
        11: "Họ và tên: {{ aftersale_customer_name }}",
        12: (
            "Số CMND/CCCD: {{ aftersale_customer_id_number }}    "
            "Ngày cấp: {{ aftersale_customer_issue_date }}    "
            "Nơi cấp: {{ aftersale_customer_issue_place }}"
        ),
        13: "Địa chỉ: {{ aftersale_customer_address }}",
        14: "Điện thoại liên hệ cần thiết: {{ aftersale_customer_phone }}",
        15: "Giấy tờ kèm theo:",
        16: "{{ id_attachment_mark }} CCCD/CMND    {{ sim_attachment_mark }} SIM gốc",
        17: "{{ other_attachment_mark }} Giấy tờ khác (nêu rõ): {{ other_attachment_value }}",
        18: "",
        22: "{{ update_information_mark }} Cập nhật thông tin",
        23: (
            "Tôi xin cam kết là chủ sở hữu của số điện thoại Vietnamobile: "
            "{{ update_subscriber_number }}. Tôi đã cung cấp cho cửa hàng SIM gốc và cam kết "
            "thuê bao không vướng bất kỳ tranh chấp nào."
        ),
        24: "{{ replace_sim_mark }} Thay SIM",
        25: (
            "Tôi xin cam kết là chủ sở hữu của số điện thoại Vietnamobile: "
            "{{ replace_sim_subscriber_number }}. Trong trường hợp xảy ra bất kỳ tranh chấp "
            "về việc thay SIM cho số thuê bao này, tôi cam đoan sẽ phối hợp với Vietnamobile "
            "để giải quyết."
        ),
        26: "{{ transfer_mark }} Chuyển chủ quyền",
        27: (
            "Tôi đồng ý thanh lý Hợp đồng cung cấp và sử dụng dịch vụ thông tin di động mặt đất "
            "Vietnamobile của thuê bao {{ transfer_subscriber_number }} và chuyển quyền sử dụng số "
            "thuê bao này và dịch vụ điện thoại di động trả trước cho Ông/Bà "
            "{{ transfer_new_owner_name }}, số CMND/CCCD {{ transfer_new_owner_id_number }}, "
            "ngày cấp {{ transfer_new_owner_issue_date }}, nơi cấp "
            "{{ transfer_new_owner_issue_place }} (“Chủ thuê bao mới”)."
        ),
        29: (
            "Tôi ({{ requester_role_mark }} Người yêu cầu hoặc {{ new_owner_role_mark }} "
            "Chủ thuê bao mới) là chủ sở hữu của số thuê bao {{ common_subscriber_number }}. "
            "Tôi đã được nhân viên tư vấn đầy đủ về các quyền và nghĩa vụ của gói cước đi kèm "
            "số thuê bao này và tôi đồng ý tiếp tục sử dụng và thực hiện các cam kết của gói "
            "cước theo quy định của Vietnamobile."
        ),
        31: (
            "Tôi đồng ý để Vietnamobile thu hồi lại số thuê bao vô điều kiện hoặc áp dụng các "
            "biện pháp khác trong trường hợp tôi vi phạm điều khoản đã cam kết hoặc có bất kỳ "
            "khiếu nại nào từ chủ thuê bao cũ và/hoặc bên thứ ba khác và Vietnamobile không liên "
            "hệ được với tôi qua số điện thoại 1: {{ backup_phone_1_line }} hoặc số điện thoại 2: "
            "{{ backup_phone_2_line }} trong vòng 24 giờ. Tôi cam đoan sẽ phối hợp với Vietnamobile "
            "để giải quyết và chấp nhận quyết định cuối cùng của Vietnamobile."
        ),
    }
    for index, value in values.items():
        set_text(paragraphs[index], value)

    # Repair the last signature area against the original form.  The edited
    # file accidentally joined "GIAO DỊCH VIÊN" to the previous caption and
    # introduced a redundant two-column section.  The legal form has exactly
    # three independent signature columns.
    if len(paragraphs) >= 39 and "GIAO DỊCH VIÊN" in paragraphs[36].text:
        set_text(paragraphs[33], "NGƯỜI YÊU CẦU")
        set_text(paragraphs[34], "(Ký và ghi rõ họ tên)")
        set_text_after_column_break(paragraphs[35], "CHỦ THUÊ BAO MỚI")
        set_text(paragraphs[36], "(Ký và ghi rõ họ tên)")

        dealer_heading_xml = deepcopy(paragraphs[35]._p)
        paragraphs[36]._p.addnext(dealer_heading_xml)
        dealer_heading = Paragraph(dealer_heading_xml, paragraphs[36]._parent)
        set_text_after_column_break(dealer_heading, "GIAO DỊCH VIÊN\n{{ aftersale_staff_name }}")
        set_text(paragraphs[37], "(Ký và ghi rõ họ tên)")

        # Remove only the accidental paragraph-level section after signatures.
        accidental_section = paragraphs[38]
        if accidental_section._p.xpath("./w:pPr/w:sectPr"):
            accidental_section._element.getparent().remove(accidental_section._element)
    else:
        dealer_heading = next(
            (p for p in document.paragraphs if "GIAO DỊCH VIÊN" in p.text),
            None,
        )
        if dealer_heading is not None:
            set_text_after_column_break(
                dealer_heading,
                "GIAO DỊCH VIÊN\n{{ aftersale_staff_name }}",
            )

    final_section = document._element.body.sectPr
    final_columns = final_section.find(qn("w:cols"))
    if final_columns is None:
        final_columns = OxmlElement("w:cols")
        document_grid = final_section.find(qn("w:docGrid"))
        if document_grid is None:
            final_section.append(final_columns)
        else:
            document_grid.addprevious(final_columns)
    final_columns.set(qn("w:num"), "3")

    # Word's original list numbering draws its own empty square.  The selected
    # square is data-driven now, so keeping numPr would render two checkboxes.
    for index in (22, 24, 26):
        properties = paragraphs[index]._p.pPr
        if properties is not None and properties.numPr is not None:
            properties.remove(properties.numPr)

    # Do not normalize fonts/spacing here.  The canonical document is maintained
    # visually in Word/LibreOffice, so migration must preserve that design.
    document.save(path)


def migrate_prepaid_contract() -> None:
    source = ROOT / (
        "desktop_app/data/source/original_documents/tra truoc/"
        "0.1. Trả trước_Hợp đồng theo mẫu trả trước (Sep 2025).docx"
    )
    destination = ROOT / (
        "desktop_app/backend/documents/prepaid_contract/"
        "00_MAU_HOP_DONG_TRA_TRUOC.docx"
    )
    document = Document(source)
    paragraphs = document.paragraphs
    values = {
        5: "Hợp đồng số: {{ contract_number }}    Mã thuê bao: {{ subscriber_code }}",
        12: "{{ document_date_line }}",
        14: "Tên cơ quan/tổ chức (Viết in hoa): {{ prepaid_organization_name }}",
        15: "Địa chỉ trụ sở chính: {{ prepaid_headquarters_address }}",
        16: (
            "Số QĐTL/GCNĐKKD&ĐKĐT/GPĐT/GCNĐKDN: {{ prepaid_business_number }}    "
            "Nơi cấp: {{ prepaid_business_issue_place }}    Ngày cấp: {{ prepaid_business_issue_date }}"
        ),
        17: (
            "Người đại diện/ủy quyền: {{ prepaid_representative_name }}    "
            "Chức vụ: {{ prepaid_representative_position }}    "
            "Giấy ủy quyền: {{ prepaid_authorization }}"
        ),
        18: (
            "Số định danh cá nhân/Số định danh điện tử/Hộ chiếu: {{ prepaid_organization_id_number }}    "
            "Nơi cấp: {{ prepaid_organization_issue_place }}    "
            "Ngày cấp: {{ prepaid_organization_issue_date }}"
        ),
        19: "Ngày, tháng, năm sinh: {{ prepaid_organization_birth_date }}",
        20: (
            "Điện thoại liên hệ: {{ prepaid_organization_phone }}    "
            "Email: {{ prepaid_organization_email }}    Liên hệ khác: {{ prepaid_organization_other_contact }}"
        ),
        22: "Họ tên (Viết in hoa): {{ prepaid_individual_name }}",
        23: (
            "Số định danh cá nhân/Số định danh điện tử/Hộ chiếu: {{ prepaid_individual_id_number }}    "
            "Nơi cấp: {{ prepaid_individual_issue_place }}    Ngày cấp: {{ prepaid_individual_issue_date }}"
        ),
        24: "Ngày, tháng, năm sinh: {{ prepaid_individual_birth_date }}",
        25: "Địa chỉ theo CCCD/Căn cước/Hộ chiếu: {{ prepaid_individual_address }}",
        26: "",
        27: (
            "Điện thoại liên hệ: {{ prepaid_individual_phone }}    "
            "Email: {{ prepaid_individual_email }}    Liên hệ khác: {{ prepaid_individual_other_contact }}"
        ),
        28: "Quốc tịch: {{ prepaid_individual_nationality }}",
        32: "Địa chỉ: {{ shop_address }}",
        33: "Người đại diện: {{ provider_representative }}    Chức vụ: {{ provider_position }}",
        37: "Điểm cung cấp dịch vụ viễn thông: {{ service_point_name }}",
        38: "Họ tên nhân viên giao dịch: {{ staff_name }}",
        44: "Địa chỉ điểm giao dịch: {{ shop_address }}",
        45: "Số điện thoại của điểm giao dịch: {{ shop_phone }}",
        46: (
            "Thời gian thực hiện đăng ký thông tin thuê bao: {{ registration_time }}. "
            "Bên A và Bên B đồng ý giao kết hợp đồng cung cấp và sử dụng dịch vụ thông tin di động "
            "mặt đất Vietnamobile (hình thức thanh toán trả trước) (“Hợp đồng”) với các nội dung sau đây:"
        ),
    }
    for index, value in values.items():
        set_text(paragraphs[index], value)

    subscriber_table = document.tables[0]
    for cell, placeholder in zip(
        subscriber_table.rows[1].cells,
        ("subscriber_number", "sim_serial", "activation_date"),
    ):
        set_text(cell.paragraphs[0], "{{ " + placeholder + " }}")

    document.save(destination)


if __name__ == "__main__":
    migrate_transfer()
    migrate_aftersale()
    migrate_prepaid_contract()
    print("Đã chuyển 3 mẫu DOCX sang placeholder.")
