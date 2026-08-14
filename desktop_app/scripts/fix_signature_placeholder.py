"""One-time, idempotent fix: the "Bên B" signature name was hardcoded as
literal text ("VÕ DUY NHẬT") inside Word text boxes in the transfer and
prepaid_contract templates. renderer.py's placeholder substitution didn't
walk into text boxes at all, so this text could never be replaced — the
printed document always showed the same fixed name regardless of who
actually signs, and for prepaid_contract it could even mismatch the
{{ provider_representative }} value already used in the contract body.

Run from the repository root with the desktop virtual environment.
"""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[2]
HARDCODED_NAME = "VÕ DUY NHẬT"
PLACEHOLDER = "{{ provider_representative }}"


def _textbox_paragraphs(document: Document):
    for paragraph in document.paragraphs:
        for txbx_content in paragraph._p.findall(".//" + qn("w:txbxContent")):
            for p_element in txbx_content.findall(qn("w:p")):
                yield Paragraph(p_element, paragraph._parent)


def _set_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def fix(path: Path) -> int:
    document = Document(path)
    replaced = 0
    for paragraph in _textbox_paragraphs(document):
        if paragraph.text.strip() == HARDCODED_NAME:
            _set_text(paragraph, PLACEHOLDER)
            replaced += 1
    if replaced:
        document.save(path)
    return replaced


def main() -> None:
    targets = [
        ROOT / "desktop_app/backend/documents/transfer/00_MAU_BIEN_BAN_CHUYEN_CHU_QUYEN.docx",
        ROOT / "desktop_app/backend/documents/prepaid_contract/00_MAU_HOP_DONG_TRA_TRUOC.docx",
    ]
    for path in targets:
        count = fix(path)
        print(f"{path.name}: replaced {count} occurrence(s)")


if __name__ == "__main__":
    main()
