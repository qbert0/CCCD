"""Surgical, one-time patch to the service_registration template AFTER the
user's own direct hand-edit in Word -- per direct instruction, adds back
what that edit was still missing relative to the source PDF:

1. A "Người đại diện: ... Chức vụ: ..." line in section II.1 (Bên B's own
   signing representative -- reuses provider_representative/provider_position,
   the SAME fields prepaid_contract's own identical "Người đại diện: {{
   provider_representative }}  Chức vụ: {{ provider_position }}" line
   already uses for this exact role -- not a new concept).
2. Two real footnotes (w:footnoteReference, via convert_footnotes_to_real.py's
   own helpers -- same mechanism build_sim_change_template.py uses): one on
   "Địa chỉ theo CCCD/Căn cước/Hộ chiếu" (source PDF's own footnote 1) and
   one on the "IV. Dịch vụ cung cấp" heading (source PDF's own footnote 3,
   renumbered 2 here since this individual-only template never carries the
   org-identity footnote 2 the source PDF also has).

Deliberately NOT a rebuild via build_service_registration_template.py --
the user edited the shipped .docx directly in Word/LibreOffice after the
last generation, so only these 2 additions are made, in place, on top of
whatever is there now; nothing else in the file is touched.

Run once, from the repository root, with the desktop virtual environment:
    python desktop_app/scripts/patch_service_registration_missing_fields.py
"""

from __future__ import annotations

import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from docx.text.run import Run

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from desktop_app.backend.documents.renderer import VALUE_FONT_NAME  # noqa: E402
from desktop_app.scripts.convert_footnotes_to_real import (  # noqa: E402
    _footnote_reference_run,
    _inject_footnotes_part,
)

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "desktop_app/backend/documents/service_registration/00_MAU_PHIEU_DANG_KY_DICH_VU.docx"

FOOTNOTE_ADDRESS = "Địa chỉ trên giấy tờ dùng để đăng ký"
FOOTNOTE_SERVICES = (
    "Các nội dung bỏ trống tại phần này do hai bên thỏa thuận điền cụ thể khi kí kết "
    "Hợp đồng và phù hợp với quy định của pháp luật"
)


def _insert_paragraph_after(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def _format_run(run, *, size=12, bold=False, color="000000") -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    r_pr = run._r.get_or_add_rPr()
    fonts = r_pr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), "Times New Roman")
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color)
    r_pr.append(color_node)


def _format_value_run(run) -> None:
    """Matches prepaid_contract's own already-shipped "Người đại diện:
    {{ provider_representative }}  Chức vụ: {{ provider_position }}" line
    exactly (same 2 field names, same role) -- NOT run through
    bake_field_highlighting.py's own general dots-adding pass, which
    would add "....." around these 2 tokens; that sibling line has none
    (confirmed by reading its own shipped runs directly), so this applies
    apply_value_font.py's own end-state (JetBrainsMono ExtraBold, 12pt,
    not bold) straight away instead."""
    run.font.name = VALUE_FONT_NAME
    run.font.size = Pt(12)
    run.font.bold = False
    r_pr = run._r.get_or_add_rPr()
    fonts = r_pr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), VALUE_FONT_NAME)


def patch() -> None:
    document = Document(str(TEMPLATE))

    gcn_paragraph = next(
        p for p in document.paragraphs if "Giấy chứng nhận đăng ký doanh nghiệp" in p.text
    )
    representative_paragraph = _insert_paragraph_after(gcn_paragraph)
    representative_paragraph.paragraph_format.space_after = gcn_paragraph.paragraph_format.space_after
    label1 = representative_paragraph.add_run("Người đại diện: ")
    _format_run(label1)
    token1 = representative_paragraph.add_run("{{provider_representative}}")
    _format_run(token1)
    _format_value_run(token1)
    label2 = representative_paragraph.add_run("  Chức vụ: ")
    _format_run(label2)
    token2 = representative_paragraph.add_run("{{provider_position}}")
    _format_run(token2)
    _format_value_run(token2)

    address_paragraph = next(
        p for p in document.paragraphs if p.text.startswith("Địa chỉ theo CCCD/Căn cước/Hộ chiếu")
    )
    # Split the label run so the mark sits right after "Hộ chiếu" and
    # before ": ", matching the source PDF's own "...Hộ chiếu1 (Số nhà...".
    label_run = address_paragraph.runs[0]
    prefix, _, suffix = label_run.text.partition("Hộ chiếu")
    label_run.text = prefix + "Hộ chiếu"
    footnote_element = _footnote_reference_run(1)
    label_run._r.addnext(footnote_element)
    if suffix:
        suffix_element = deepcopy(label_run._r)
        Run(suffix_element, address_paragraph).text = suffix
        footnote_element.addnext(suffix_element)

    services_heading = next(p for p in document.paragraphs if p.text.strip() == "IV. Dịch vụ cung cấp")
    services_heading.runs[-1]._r.addnext(_footnote_reference_run(2))

    document.save(str(TEMPLATE))
    _inject_footnotes_part(TEMPLATE, [(1, FOOTNOTE_ADDRESS), (2, FOOTNOTE_SERVICES)])
    print("added Người đại diện/Chức vụ line and footnotes 1-2")


if __name__ == "__main__":
    patch()
