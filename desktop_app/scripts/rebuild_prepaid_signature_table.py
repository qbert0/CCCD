"""One-time, idempotent fix: Prepaid Contract's "ĐẠI DIỆN BÊN A/B" signature
block was two fixed-size floating text boxes (`wp:anchor`). Inserting the
provider's stamp+signature image into Bên B's box (needed so the contract
carries a real signature, matching sim_change_form/transfer) overflowed the
box's fixed height and broke the layout, since these old floating boxes
don't grow to fit their content the way a table cell does. Replaces both
boxes with a real 2-column table in the same spot -- same fragility problem
and same fix as transfer's signature block.

Run from the repository root with the desktop virtual environment.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "desktop_app/backend/documents/prepaid_contract/00_MAU_HOP_DONG_TRA_TRUOC.docx"
PROVIDER_SIGNATURE = ROOT / "desktop_app/data/source/signatures/vietnamobile_provider_stamp_signature.png"

COLUMNS = [
    ("ĐẠI DIỆN BÊN A", "{{ provider_representative }}", False),
    ("ĐẠI DIỆN BÊN B", "{{ provider_representative }}", True),
]
HINT = "(Ký, ghi rõ họ tên, đóng dấu)"


def _set_cell_margins(cell, top=80, bottom=80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("bottom", bottom)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def _write(paragraph, text: str, *, size=10, bold=False, italic=False) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def _signature_cell(cell, heading: str, content: str, *, provider: bool) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    _set_cell_margins(cell)
    _write(cell.paragraphs[0], heading, bold=True)
    hint = cell.add_paragraph()
    _write(hint, HINT, size=8.5, italic=True)
    body = cell.add_paragraph()
    body.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if provider and PROVIDER_SIGNATURE.is_file():
        body.add_run().add_picture(str(PROVIDER_SIGNATURE), width=Mm(45))
        body.add_run().add_break()
    run = body.add_run(content)
    run.font.size = Pt(9.5)
    run.font.bold = True


def _is_group_anchor_host(paragraph) -> bool:
    for anchor in paragraph._p.findall(".//" + qn("wp:anchor")):
        doc_pr = anchor.find(".//" + qn("wp:docPr"))
        if doc_pr is not None and doc_pr.get("name", "").startswith("Group"):
            return True
    return False


def build(template: Path = TEMPLATE) -> int:
    document = Document(str(template))
    hosts = [p for p in document.paragraphs if _is_group_anchor_host(p)]
    if len(hosts) != 2:
        return 0

    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, (heading, content, provider) in zip(table.rows[0].cells, COLUMNS):
        _signature_cell(cell, heading, content, provider=provider)

    hosts[0]._p.addnext(table._tbl)
    for host in hosts:
        host._p.getparent().remove(host._p)
    document.save(template)
    return 1


if __name__ == "__main__":
    print(f"replaced signature block: {bool(build())}")
