"""Build the "Phiếu đăng ký dịch vụ và Bản xác nhận thông tin thuê bao
Vietnamobile trả trước" form as a real, editable Word document.

Source reference: doc/quanha/0. Trả trước_Phiếu đăng ký dịch vụ và bản
xác nhận thông tin thuê bao trả trước (Sep 2025).pdf (visual/legal
reference only, per direct instruction -- every heading, paragraph,
checkbox, line, border and answer field below is native Word content,
same approach as build_sim_change_template.py/build_ownership_confirmation_template.py,
reusing those scripts' own authoring helpers rather than re-implementing
them. Section I.2 ("Khách hàng là tổ chức") and III.2/III.3 (org
subscriber list / 4th+ subscriber) are intentionally omitted -- this
template is only ever used by QUANG_HA_STT, whose customer is always an
individual (see ServiceRegistrationSchema's own note), and the 3-row
subscriber table below IS the "03 số thuê bao đầu tiên" case the source
form itself is scoped to.

Run from the repository root with the desktop virtual environment:
    python desktop_app/scripts/build_service_registration_template.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from desktop_app.backend.documents.font_embed import embed_signature_font  # noqa: E402
from desktop_app.scripts.build_sim_change_template import (  # noqa: E402
    BLACK,
    ORANGE,
    _add_header,
    _add_line,
    _format_paragraph,
    _format_run,
    _keep,
    _section_heading,
    _set_cant_split,
    _set_cell_margins,
    _set_cell_width,
    _set_footer,
    _set_table_borders,
    _signature_cell,
    _write,
)

ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DOCX = (
    ROOT / "desktop_app" / "backend" / "documents" / "service_registration"
    / "00_MAU_PHIEU_DANG_KY_DICH_VU.docx"
)

# Section IV -- "Dịch vụ cung cấp" -- static/unticked checkbox rows, not
# data-driven. Verbatim from the source form's own service names.
DEFAULT_SERVICES = (
    "Thoại, bao gồm Hiển thị số thuê bao chủ gọi, Chuyển tiếp cuộc gọi, Chờ cuộc gọi, Cuộc gọi hội nghị",
    "Nhắn tin ngắn trong nước", "Nhắn tin ngắn quốc tế", "Gọi khẩn cấp",
    "Nhắn tin đa phương tiện", "GPRS",
)
OPTIONAL_SERVICES = (
    "In bảng kê chi tiết cước", "Gọi quốc tế", "Chuyển vùng quốc tế",
    "Hộp thư thoại", "Nhạc chuông chờ", "Dữ liệu linh hoạt",
)

# Section V -- "Điều khoản chung" -- 3 static legal paragraphs, verbatim
# from the source PDF (pdftotext -layout, cross-checked against the page
# images).
GENERAL_TERMS = [
    "Khách hàng đồng ý rằng Vietnamobile có quyền (i) thu thập và lưu giữ thông tin cá nhân, tổ chức của "
    "Khách hàng (“Thông tin thuê bao”) trong Phiếu đăng ký dịch vụ và Bản xác nhận thông tin thuê bao "
    "Vietnamobile trả trước này và các tài liệu liên quan kèm theo được quy định tại Nghị định 163/2024/NĐ-CP "
    "và các văn bản sửa đổi, bổ sung và (ii) sử dụng Thông tin thuê bao phục vụ (a) cung cấp dịch vụ viễn "
    "thông của Vietnamobile, (b) công tác bảo đảm an ninh quốc gia, trật tự an toàn xã hội, và (c) công tác "
    "quản lý nhà nước về viễn thông. Nguyên tắc, phạm vi, mục đích xử lý thông tin của Khách hàng và quyền "
    "của Khách hàng được quy định cụ thể tại Mẫu đồng ý xử lý dữ liệu cá nhân đính kèm.",
    "Khách hàng có quyền yêu cầu cập nhật, sửa đổi, hủy bỏ Thông tin thuê bao hoặc ngừng cung cấp Thông tin "
    "thuê bao của Khách hàng cho bên thứ ba, trừ trường hợp việc cung cấp thông tin là theo yêu cầu của cơ "
    "quan nhà nước có thẩm quyền, tại bất kỳ điểm cung cấp dịch vụ viễn thông hoặc điểm cung cấp dịch vụ "
    "viễn thông được ủy quyền nào của Vietnamobile.",
    "Bản Điều kiện giao dịch chung của dịch vụ thông tin di động mặt đất (hình thức thanh toán trả trước) "
    "(“Điều kiện chung”) được niêm yết tại các điểm cung cấp dịch vụ viễn thông, trên website của Bên B "
    "và cung cấp cho Bên A thông qua các hình thức như bản in trực tiếp, qua email hoặc các phương thức "
    "khác do hai bên thỏa thuận. Khách hàng và Vietnamobile cam kết tuân thủ các quy định tại Điều kiện chung.",
]


def _build_section_one(document: Document) -> None:
    title = document.add_paragraph()
    _format_paragraph(title, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    _format_run(
        title.add_run(
            "PHIẾU ĐĂNG KÝ DỊCH VỤ VÀ\nBẢN XÁC NHẬN THÔNG TIN THUÊ BAO VIETNAMOBILE TRẢ TRƯỚC"
        ),
        size=12.5, bold=True, color=ORANGE,
    )
    _keep(title, with_next=True)
    subtitle = document.add_paragraph()
    _format_paragraph(subtitle, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    _format_run(
        subtitle.add_run("(Áp dụng trong trường hợp Khách hàng đăng ký 03 số thuê bao đầu tiên)"),
        size=9.5, italic=True,
    )
    _keep(subtitle, with_next=True)
    date_line = _add_line(
        document, "Ngày ", "{{service_registration_day}} tháng {{service_registration_month}} năm {{service_registration_year}}",
        size=9.5, after=4,
    )
    date_line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _keep(date_line, with_next=True)

    section_one = _section_heading(document, "I", "Bên sử dụng dịch vụ (Gọi tắt là “Khách hàng” hoặc “Bên A”)")
    _keep(section_one, with_next=True)
    intro = document.add_paragraph()
    _format_paragraph(intro, after=2)
    _format_run(intro.add_run("1. Thông tin Khách hàng (cá nhân):"), size=9.5, bold=True)
    _keep(intro, with_next=True)

    _add_line(document, "Họ tên (Viết in hoa): ", "{{service_registration_customer_name}}", after=2)
    _add_line(
        document, "Số định danh cá nhân/Số định danh điện tử/Hộ chiếu: ",
        "{{service_registration_customer_id_number}}    Ngày cấp: {{service_registration_customer_issue_date}}"
        "    Nơi cấp: {{service_registration_customer_issue_place}}",
        after=2,
    )
    _add_line(document, "Ngày, tháng, năm sinh: ", "{{service_registration_customer_birth_date}}", after=2)
    _add_line(
        document, "Địa chỉ theo CCCD/Căn cước/Hộ chiếu: ", "{{service_registration_customer_address}}", after=2,
    )
    _add_line(
        document, "Điện thoại liên hệ: ",
        "{{service_registration_customer_phone}}    Email: {{service_registration_customer_email}}",
        after=2,
    )
    attach = _add_line(
        document, "Quốc tịch: ", "{{service_registration_customer_nationality}}", after=6,
    )
    _keep(attach, with_next=True)


def _build_section_two(document: Document) -> None:
    section_two = _section_heading(document, "II", "Bên cung cấp dịch vụ (Gọi tắt là “Vietnamobile” hoặc “Bên B”)")
    _keep(section_two, with_next=True)

    provider = document.add_paragraph()
    _format_paragraph(provider, after=1)
    _format_run(
        provider.add_run(
            "1. Đơn vị cung cấp dịch vụ viễn thông: Công ty Cổ phần Viễn thông Di động Vietnamobile"
        ),
        size=9.5, bold=True,
    )
    _keep(provider, with_next=True)
    for line in (
        "Giấy chứng nhận đăng ký doanh nghiệp: 0107429715 do Sở Kế hoạch và Đầu tư thành phố Hà Nội "
        "(nay là Sở Tài chính thành phố Hà Nội) cấp lần đầu ngày 12/5/2016",
        "Điện thoại: (024) 35730123",
        "Thư điện tử: cskh@vietnamobile.com.vn        Website: http://www.vietnamobile.com.vn",
        "Mã số thuế: 0107429715",
    ):
        p = document.add_paragraph()
        _format_paragraph(p, after=1)
        _format_run(p.add_run(line), size=9)
        _keep(p, with_next=True)

    service_point = document.add_paragraph()
    _format_paragraph(service_point, before=3, after=2)
    _format_run(service_point.add_run("2. Điểm cung cấp dịch vụ viễn thông: {{service_point_name}}"), size=9.5, bold=True)
    _keep(service_point, with_next=True)
    _add_line(document, "Họ tên nhân viên giao dịch: ", "{{staff_name}}", after=2)
    _add_line(document, "Địa chỉ điểm giao dịch: ", "{{service_point_address}}", after=2)
    _add_line(document, "Số điện thoại của điểm giao dịch: ", "{{service_point_phone}}", after=2)
    reg_time = _add_line(
        document, "Thời gian thực hiện đăng ký thông tin thuê bao: ", "{{registration_time}}", after=6,
    )
    _keep(reg_time, with_next=True)


def _build_subscriber_table(document: Document) -> None:
    section_three = _section_heading(document, "III", "Thông tin cá nhân sử dụng số thuê bao")
    _keep(section_three, with_next=True)
    intro = document.add_paragraph()
    _format_paragraph(intro, after=2)
    _format_run(intro.add_run("Danh sách số thuê bao đăng ký của Khách hàng:"), size=9.5)
    _keep(intro, with_next=True)

    table = document.add_table(rows=4, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    widths = (14, 55, 55, 55)
    headers = ("TT", "Số thuê bao", "Số sê-ri SIM", "Ngày hòa mạng")
    for cell, width, text in zip(table.rows[0].cells, widths, headers):
        _set_cell_width(cell, width)
        _write(cell, text, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    for index, row in enumerate(table.rows[1:], start=1):
        for cell, width in zip(row.cells, widths):
            _set_cell_width(cell, width)
            _set_cell_margins(cell)
        _write(row.cells[0], str(index), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cant_split(row)
    _set_cant_split(table.rows[0])


def _checkbox_row(table, left: str, right: str = "") -> None:
    cells = table.add_row().cells
    _write(cells[0], "☐", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    _write(cells[1], left, size=8.7)
    if right:
        _write(cells[2], "☐", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        _write(cells[3], right, size=8.7)
    else:
        cells[1].merge(cells[3])
    table.rows[-1].height = Mm(9)
    table.rows[-1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    _set_cant_split(table.rows[-1])


def _build_services_section(document: Document) -> None:
    section_four = _section_heading(document, "IV", "Dịch vụ cung cấp")
    _keep(section_four, with_next=True)
    default_label = document.add_paragraph()
    _format_paragraph(default_label, after=1)
    _format_run(default_label.add_run("1. Các dịch vụ mặc định:"), size=9.5, bold=True)
    _keep(default_label, with_next=True)

    table = document.add_table(rows=0, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    widths = (11, 84, 11, 84)
    services = list(DEFAULT_SERVICES)
    for i in range(0, len(services) - 1, 2):
        row_cells = table.add_row().cells
        for cell, width in zip(row_cells, widths):
            _set_cell_width(cell, width)
        _write(row_cells[0], "☐", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        _write(row_cells[1], services[i], size=8.5)
        _write(row_cells[2], "☐", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        _write(row_cells[3], services[i + 1], size=8.5)
        table.rows[-1].height = Mm(9)
        table.rows[-1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        _set_cant_split(table.rows[-1])
    if len(services) % 2:
        row_cells = table.add_row().cells
        for cell, width in zip(row_cells, widths):
            _set_cell_width(cell, width)
        _write(row_cells[0], "☐", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        merged = row_cells[1].merge(row_cells[3])
        _write(merged, services[-1], size=8.5)
        _set_cant_split(table.rows[-1])

    optional_label = document.add_paragraph()
    _format_paragraph(optional_label, before=3, after=1)
    _format_run(optional_label.add_run("2. Các dịch vụ đăng ký:"), size=9.5, bold=True)
    _keep(optional_label, with_next=True)

    table2 = document.add_table(rows=0, cols=4)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    table2.autofit = False
    _set_table_borders(table2)
    for i in range(0, len(OPTIONAL_SERVICES), 2):
        pair = OPTIONAL_SERVICES[i:i + 2]
        row_cells = table2.add_row().cells
        for cell, width in zip(row_cells, widths):
            _set_cell_width(cell, width)
        _write(row_cells[0], "☐", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        _write(row_cells[1], pair[0], size=8.5)
        if len(pair) > 1:
            _write(row_cells[2], "☐", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
            _write(row_cells[3], pair[1], size=8.5)
        else:
            row_cells[1].merge(row_cells[3])
        table2.rows[-1].height = Mm(9)
        table2.rows[-1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        _set_cant_split(table2.rows[-1])


def _build_terms_section(document: Document) -> None:
    section_five = _section_heading(document, "V", "Điều khoản chung")
    _keep(section_five, with_next=True)
    for index, text in enumerate(GENERAL_TERMS, start=1):
        paragraph = document.add_paragraph()
        _format_paragraph(paragraph, after=3)
        _format_run(paragraph.add_run(f"{index}. {text}"), size=9)


def _build_signature_block(document: Document) -> None:
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for cell in table.rows[0].cells:
        _set_cell_width(cell, 63)
    table.rows[0].height = Mm(45)
    table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    _set_cant_split(table.rows[0])
    _signature_cell(
        table.cell(0, 0), "KHÁCH HÀNG",
        given_key="customer_signature_given_name", full_key="customer_signature_name",
    )
    _signature_cell(
        table.cell(0, 1), "ĐẠI DIỆN BÊN CUNG CẤP\nDỊCH VỤ VIỄN THÔNG",
        given_key="provider_representative_signature_given_name", full_key="provider_representative_signature",
    )
    _signature_cell(
        table.cell(0, 2), "GIAO DỊCH VIÊN",
        given_key="service_registration_clerk_signature_given_name",
        full_key="service_registration_clerk_signature_name",
    )


def build(output_docx: Path = OUTPUT_DOCX) -> Path:
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(15)
    section.bottom_margin = Mm(14)
    section.left_margin = Mm(16)
    section.right_margin = Mm(16)
    section.header_distance = Mm(2)
    section.footer_distance = Mm(2)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(9.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    embed_signature_font(document)
    _add_header(section)
    _set_footer(section)
    _build_section_one(document)
    _build_section_two(document)
    _build_subscriber_table(document)
    _build_services_section(document)
    _build_terms_section(document)
    _build_signature_block(document)

    document.core_properties.title = "Phiếu đăng ký dịch vụ và Bản xác nhận thông tin thuê bao Vietnamobile trả trước"
    document.core_properties.subject = "Biểu mẫu Word native dựng lại từ PDF tham khảo (doc/quanha)"
    document.core_properties.comments = "Không sử dụng ảnh nền hoặc ảnh chụp trang PDF."
    document.save(output_docx)
    return output_docx


if __name__ == "__main__":
    print(build())
