"""Build the ownership-confirmation form ("Giấy cam kết - Dành cho Giao
dịch xác nhận quyền") as a real, editable Word document.

Source reference: doc/quanha/Form - Xác nhận quyền.pdf (visual/legal
reference only, per direct instruction -- every heading, paragraph,
checkbox, line, border and answer field below is native Word content,
same approach as build_sim_change_template.py, reusing that script's own
authoring helpers rather than re-implementing them.

Run from the repository root with the desktop virtual environment:
    python desktop_app/scripts/build_ownership_confirmation_template.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from desktop_app.backend.documents.font_embed import embed_signature_font  # noqa: E402
from desktop_app.scripts.build_sim_change_template import (  # noqa: E402
    BLACK,
    ORANGE,
    _add_header,
    _add_line,
    _format_paragraph,
    _format_run,
    _keep,
    _section_heading,
    _set_cant_split,
    _set_cell_margins,
    _set_cell_width,
    _set_footer,
    _set_table_borders,
    _signature_cell,
    _write,
)

ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DOCX = (
    ROOT / "desktop_app" / "backend" / "documents" / "ownership_confirmation"
    / "00_MAU_GIAY_CAM_KET_XAC_NHAN_QUYEN.docx"
)

COMMITMENTS = [
    "Tôi xin cam kết là tôi là người đã đăng ký thông tin và sử dụng, có quyền sử dụng đối với "
    "(các) số thuê bao Vietnamobile như đã nêu trên.",
    "Tôi đã cung cấp cho cửa hàng SIM gốc và cam kết thuê bao không có bất kỳ tranh chấp nào.",
    "Tôi đã được nhân viên tư vấn đầy đủ về các quyền và nghĩa vụ của gói cước đi kèm (các) số "
    "thuê bao này và tôi đồng ý tiếp tục sử dụng và thực hiện các cam kết của gói cước theo quy "
    "định của Vietnamobile.",
    "Trường hợp có xảy ra khiếu kiện, tranh chấp đối với (các) số thuê bao này và/hoặc dịch vụ "
    "yêu cầu cho (các) số thuê bao này (Cập nhật lại thông tin/Thay Sim/Chuyển chủ quyền) hoặc có "
    "bất cứ vấn đề liên quan đến (các) số thuê bao này, tôi cam đoan sẽ phối hợp với Vietnamobile "
    "để giải quyết và chấp nhận quyết định cuối cùng của Vietnamobile và Vietnamobile có quyền thu "
    "hồi lại (các) số thuê bao này theo quy định của pháp luật. Tôi đồng ý không chuyển mạng giữ số "
    "trong vòng 36 tháng kể từ ngày ký Giấy Cam kết này để Vietnamobile có cơ sở giải quyết yêu cầu "
    "liên quan đến (các) số thuê bao này.",
    "Tôi đồng ý để Vietnamobile thu hồi lại số thuê bao vô điều kiện hoặc áp dụng các biện pháp "
    "khác trong trường hợp tôi vi phạm điều khoản đã cam kết hoặc có bất kỳ khiếu nại nào từ chủ "
    "thuê bao cũ và/hoặc bên thứ ba khác và Vietnamobile không liên hệ được với tôi trong vòng 24 "
    "giờ. Tôi cam đoan sẽ phối hợp với Vietnamobile để giải quyết và chấp nhận quyết định cuối "
    "cùng của Vietnamobile.",
    "Tôi xin chịu trách nhiệm trước pháp luật và Công ty Cổ phần Viễn thông Di động Vietnamobile "
    "về Giấy Cam kết và các nội dung thông tin cung cấp tại Giấy Cam kết này.",
]


def _build_body(document: Document) -> None:
    shop_line = document.add_paragraph()
    _format_paragraph(shop_line, align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    _format_run(shop_line.add_run("Cửa hàng: ..............................................."), size=9)
    _keep(shop_line, with_next=True)
    address_line = document.add_paragraph()
    _format_paragraph(address_line, align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    _format_run(address_line.add_run("Địa chỉ: ..............................................."), size=9)
    _keep(address_line, with_next=True)
    phone_line = document.add_paragraph()
    _format_paragraph(phone_line, align=WD_ALIGN_PARAGRAPH.RIGHT, after=6)
    _format_run(phone_line.add_run("Điện thoại: ..............................................."), size=9)
    _keep(phone_line, with_next=True)

    title = document.add_paragraph()
    _format_paragraph(title, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    _format_run(title.add_run("GIẤY CAM KẾT"), size=14, bold=True, color=ORANGE)
    _keep(title, with_next=True)
    subtitle = document.add_paragraph()
    _format_paragraph(subtitle, align=WD_ALIGN_PARAGRAPH.CENTER, after=5)
    _format_run(subtitle.add_run("(Dành cho Giao dịch xác nhận quyền)"), size=10.5, italic=True)
    _keep(subtitle, with_next=True)

    date_line = _add_line(
        document, "Ngày ", "{{ownership_day}} tháng {{ownership_month}} năm {{ownership_year}}",
        size=9.5, after=4,
    )
    _keep(date_line, with_next=True)

    _add_line(document, "Họ & Tên: ", "{{ownership_customer_name}}", after=2)
    _add_line(
        document, "Số CMND/CCCD: ",
        "{{ownership_customer_id_number}}    Ngày cấp: {{ownership_customer_issue_date}}"
        "    Nơi cấp: {{ownership_customer_issue_place}}",
        after=2,
    )
    _add_line(document, "Địa chỉ: ", "{{ownership_customer_address}}", after=2)
    attach_line = _add_line(
        document, "Điện thoại liên hệ cần thiết: ", "{{ownership_customer_phone}}", after=2,
    )
    _keep(attach_line, with_next=True)

    attach_label = document.add_paragraph()
    _format_paragraph(attach_label, after=1)
    _format_run(attach_label.add_run("Giấy tờ kèm theo:"), size=9.5)
    _keep(attach_label, with_next=True)
    attach_row = document.add_paragraph()
    _format_paragraph(attach_row, after=1)
    _format_run(
        attach_row.add_run(
            "{{id_attachment_mark}} CMND/CCCD          {{sim_attachment_mark}} Sim gốc          "
            "{{other_attachment_mark}} Giấy tờ khác (ghi rõ): {{other_attachment_value}}"
        ),
        size=9.5,
    )
    _keep(attach_row, with_next=True)
    note = document.add_paragraph()
    _format_paragraph(note, after=6)
    _format_run(note.add_run('(Sau đây gọi là "Khách hàng" hoặc "Người yêu cầu")'), size=9, italic=True)
    _keep(note, with_next=True)

    section_one = _section_heading(document, "1", "Phần Khách hàng yêu cầu")
    _keep(section_one, with_next=True)
    service_line = document.add_paragraph()
    _format_paragraph(service_line, after=3)
    _format_run(
        service_line.add_run("Dịch vụ yêu cầu: Xác nhận quyền đối với số thuê bao Vietnamobile"),
        size=9.5, bold=True,
    )
    _keep(service_line, with_next=True)

    table = document.add_table(rows=2, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    widths = (14, 40, 68, 34, 24)
    headers = ("STT", "Số thuê bao", "Họ và tên cá nhân sử dụng số thuê bao", "Số GTTT", "Ngày cấp")
    for cell, width, text in zip(table.rows[0].cells, widths, headers):
        _set_cell_width(cell, width)
        _write(cell, text, bold=True, size=8.7, align=WD_ALIGN_PARAGRAPH.CENTER)
    for cell, width in zip(table.rows[1].cells, widths):
        _set_cell_width(cell, width)
        _set_cell_margins(cell)
    for row in table.rows:
        _set_cant_split(row)

    section_two = _section_heading(document, "2", "Cam kết của Người yêu cầu")
    _keep(section_two, with_next=True)
    for index, text in enumerate(COMMITMENTS, start=1):
        paragraph = document.add_paragraph()
        _format_paragraph(paragraph, after=3)
        _format_run(paragraph.add_run(f"2.{index}. {text}"), size=9)


def _build_signature_block(document: Document) -> None:
    date = document.add_paragraph()
    _format_paragraph(date, align=WD_ALIGN_PARAGRAPH.RIGHT, before=6, after=1)
    date.add_run().add_break()
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for cell in table.rows[0].cells:
        _set_cell_width(cell, 95)
    table.rows[0].height = Mm(45)
    table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    _set_cant_split(table.rows[0])
    _signature_cell(
        table.cell(0, 0), "NGƯỜI YÊU CẦU",
        given_key="ownership_requester_signature_given_name",
        full_key="ownership_requester_signature_name",
    )
    _signature_cell(
        table.cell(0, 1), "GIAO DỊCH VIÊN",
        given_key="ownership_clerk_signature_given_name",
        full_key="ownership_clerk_signature_name",
    )


def build(output_docx: Path = OUTPUT_DOCX) -> Path:
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(17)
    section.bottom_margin = Mm(16)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    section.header_distance = Mm(2)
    section.footer_distance = Mm(2)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(9.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    embed_signature_font(document)
    _add_header(section)
    _set_footer(section)
    _build_body(document)
    _build_signature_block(document)

    document.core_properties.title = "Giấy cam kết - Dành cho Giao dịch xác nhận quyền"
    document.core_properties.subject = "Biểu mẫu Word native dựng lại từ PDF tham khảo (doc/quanha)"
    document.core_properties.comments = "Không sử dụng ảnh nền hoặc ảnh chụp trang PDF."
    document.save(output_docx)
    return output_docx


if __name__ == "__main__":
    print(build())
