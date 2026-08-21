"""Build the SIM-change form as a real, editable Word document.

The source PDF is used only as a visual/legal reference. Every heading,
paragraph, checkbox, line, border and answer field below is native Word
content, including every signature -- the header/footer brand banners are
the only images baked into the template at all now (see _signature_cell
and font_embed.py for why a script font replaced real signature images).

Footnotes are printed as real per-page FOOTER content (pinned to the page
bottom, above the brand banner) -- NOT real OOXML
w:footnoteReference/footnotes.xml. That was tried first and verified
broken: under this project's actual render pipeline (soffice --headless
--convert-to pdf, the same one docx_to_images.py uses to produce the
customer-facing JPGs), LibreOffice silently drops a footnote's body text
once more than one footnote crosses a hard page break, and shows the text
of an EARLIER footnote under a LATER page's reference number instead
(reproduced in isolation with a minimal 3-footnote/2-page document;
persists through an .odt round-trip, so it's a Writer layout bug, not an
export-filter quirk). A first fix printed the note text as an ordinary
paragraph right before each page's break -- correct and always on the
right page, but visually indistinguishable from body text since it just
flowed after the last table row instead of sitting at the true bottom of
the page. Each page is now its own section (page 1/2/3 need different
footer text) with an unlinked footer holding the note paragraphs above
the same footer image every page already had; footers are ordinary,
well-supported per-section content with none of the footnote bug above.

Every signature is 2 lines in the embedded "Great Vibes" script font (see
font_embed.py): "ký tên" (just the given name -- the actual signature
stroke) above "ghi rõ họ tên" (the full name).
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from desktop_app.backend.documents.font_embed import apply_signature_font, embed_signature_font  # noqa: E402


ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DOCX = ROOT / "desktop_app" / "backend" / "documents" / "sim_change_form" / "00_MAU_PHIEU_THAY_DOI_DICH_VU_TRA_TRUOC.docx"
HEADER_IMAGE = ROOT / "desktop_app" / "data" / "source" / "branding" / "vietnamobile_header.png"
FOOTER_IMAGE = ROOT / "desktop_app" / "data" / "source" / "branding" / "vietnamobile_footer.png"
ORANGE = "F45A1E"
BLACK = "000000"

FOOTNOTES = {
    1: (
        "Số Quyết định thành lập/Giấy chứng nhận đăng ký kinh doanh và đăng ký đầu tư/"
        "Giấy phép đầu tư/Giấy chứng nhận đăng ký doanh nghiệp hoặc giấy tờ chứng minh "
        "pháp nhân khác."
    ),
    2: "Số Định danh cá nhân/Định danh điện tử/Hộ chiếu.",
    3: (
        "Các nội dung bỏ trống tại Phần II, III do hai bên thỏa thuận điền cụ thể khi "
        "ký kết phiếu và phù hợp với quy định của pháp luật."
    ),
}


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def _set_cant_split(row) -> None:
    """Keep one logical table row on one page."""
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def _set_cell_margins(cell, top=55, start=65, bottom=55, end=65) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table, color=BLACK, size=5) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def _set_cell_width(cell, width_mm: float) -> None:
    cell.width = Mm(width_mm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(round(width_mm * 56.6929)))
    tc_w.set(qn("w:type"), "dxa")


def _format_run(run, *, size=9.5, bold=False, italic=False, color=BLACK) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    r_pr = run._r.get_or_add_rPr()
    fonts = r_pr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), "Times New Roman")
    color_node = r_pr.find(qn("w:color"))
    if color_node is None:
        color_node = OxmlElement("w:color")
        r_pr.append(color_node)
    color_node.set(qn("w:val"), color)


def _format_paragraph(paragraph, *, align=None, before=0, after=0, line=1.0) -> None:
    if align is not None:
        paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.widow_control = True


def _keep(paragraph, *, with_next=False, together=True) -> None:
    paragraph.paragraph_format.keep_together = together
    paragraph.paragraph_format.keep_with_next = with_next


def _add_footnote_marker(paragraph, number: int, *, size=7.5) -> None:
    """Inline superscript numeral matching a note printed by _add_page_footnotes."""
    run = paragraph.add_run(str(number))
    _format_run(run, size=size)
    vert_align = OxmlElement("w:vertAlign")
    vert_align.set(qn("w:val"), "superscript")
    run._r.get_or_add_rPr().append(vert_align)


def _set_footer(section, numbers: list[int]) -> None:
    """Build this section's own footer: the given FOOTNOTES entries (if
    any) printed above the brand banner image, always pinned to the page
    bottom -- see the module docstring for why these aren't real OOXML
    footnotes. Each page needing different footnote text must be its own
    section (see _build_page_two/_build_page_three), since a section's
    footer is otherwise linked to and identical to the previous one's."""
    footer = section.footer
    footer.is_linked_to_previous = False
    usable_mm = (section.page_width - section.left_margin - section.right_margin) / 36000

    first = footer.paragraphs[0]
    first.text = ""
    if numbers:
        _format_paragraph(first, after=2)
        borders = OxmlElement("w:pBdr")
        top = OxmlElement("w:top")
        top.set(qn("w:val"), "single")
        top.set(qn("w:sz"), "4")
        top.set(qn("w:space"), "1")
        top.set(qn("w:color"), BLACK)
        borders.append(top)
        first._p.get_or_add_pPr().append(borders)
        for number in numbers:
            p = footer.add_paragraph()
            _format_paragraph(p, after=1, line=1.0)
            number_run = p.add_run(str(number))
            _format_run(number_run, size=7.5)
            vert_align = OxmlElement("w:vertAlign")
            vert_align.set(qn("w:val"), "superscript")
            number_run._r.get_or_add_rPr().append(vert_align)
            text_run = p.add_run(f" {FOOTNOTES[number]}")
            _format_run(text_run, size=7.5, italic=True)
        image_paragraph = footer.add_paragraph()
    else:
        image_paragraph = first

    _format_paragraph(image_paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    if FOOTER_IMAGE.is_file():
        image_paragraph.add_run().add_picture(str(FOOTER_IMAGE), width=Mm(usable_mm))


def _write(cell_or_paragraph, text: str, *, size=9.5, bold=False, italic=False,
           align=None, color=BLACK, clear=True):
    if hasattr(cell_or_paragraph, "paragraphs"):
        cell = cell_or_paragraph
        if clear:
            cell.text = ""
        paragraph = cell.paragraphs[0]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_margins(cell)
    else:
        paragraph = cell_or_paragraph
        if clear:
            paragraph.clear()
    _format_paragraph(paragraph, align=align)
    run = paragraph.add_run(text)
    _format_run(run, size=size, bold=bold, italic=italic, color=color)
    return paragraph


def _write_segments(cell, segments, *, size=9.5, align=None):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    _set_cell_margins(cell)
    paragraph = cell.paragraphs[0]
    _format_paragraph(paragraph, align=align)
    for segment in segments:
        if isinstance(segment, int):
            _add_footnote_marker(paragraph, segment)
        else:
            _format_run(paragraph.add_run(segment), size=size)
    return paragraph


def _add_line(document, label: str, placeholder: str = "", *, size=9.5, bold_label=False,
              after=2, indent_mm=0):
    p = document.add_paragraph()
    _format_paragraph(p, after=after)
    if indent_mm:
        p.paragraph_format.left_indent = Mm(indent_mm)
    run = p.add_run(label)
    _format_run(run, size=size, bold=bold_label)
    if placeholder:
        value = p.add_run(placeholder)
        _format_run(value, size=size)
    return p


def _section_heading(document, number: str, title: str):
    p = document.add_paragraph()
    _format_paragraph(p, before=2, after=2)
    run = p.add_run(f"{number}. {title}")
    _format_run(run, size=10.5, bold=True, color=ORANGE)
    _keep(p, with_next=True)
    return p


def _add_header(section) -> None:
    header = section.header
    header.paragraphs[0].text = ""
    p = header.paragraphs[0]
    _format_paragraph(p, align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    if HEADER_IMAGE.is_file():
        p.add_run().add_picture(str(HEADER_IMAGE), width=Mm(63))


def _build_page_one(document: Document) -> None:
    title = document.add_paragraph()
    _format_paragraph(title, align=WD_ALIGN_PARAGRAPH.CENTER, after=5)
    _format_run(
        title.add_run("PHIẾU CUNG CẤP VÀ THAY ĐỔI DỊCH VỤ THÔNG TIN DI ĐỘNG MẶT ĐẤT\n(HÌNH THỨC THANH TOÁN TRẢ TRƯỚC)"),
        size=12,
        bold=True,
        color=ORANGE,
    )
    _keep(title, with_next=True)
    _section_heading(document, "I", "THÔNG TIN KHÁCH HÀNG")

    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    _set_cell_width(table.cell(0, 0), 95)
    _set_cell_width(table.cell(0, 1), 95)
    _write(table.cell(0, 0), "Thông tin khách hàng đã cung cấp", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _write(table.cell(0, 1), "Thông tin khách hàng thay đổi (Nếu có)", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_repeat_table_header(table.rows[0])
    _set_cant_split(table.rows[0])

    rows = [
        (
            "Tên cơ quan, tổ chức hoặc cá nhân (viết in hoa):\n{{sim_customer_name}}",
            "Tên cơ quan, tổ chức hoặc cá nhân (viết in hoa):\n................................................................",
            14,
        ),
        (
            "Địa chỉ trụ sở chính/Địa chỉ theo CCCD/Căn cước/Hộ chiếu:\n{{sim_customer_address}}",
            "Địa chỉ trụ sở chính/Địa chỉ theo CCCD/Căn cước/Hộ chiếu:\n................................................................",
            18,
        ),
        (
            [
                "- Số QĐTL/GCNĐKKD&ĐKĐT/GPĐT/GCNĐKDN",
                1,
                ":\n{{sim_customer_id_number}}\n"
                "- Nơi cấp/Đơn vị cấp: {{sim_customer_issue_place}}\n"
                "- Ngày cấp: {{sim_customer_issue_date}}",
            ],
            "- Số QĐTL/GCNĐKKD&ĐKĐT/GPĐT/GCNĐKDN:\n.................................................................\n"
            "- Nơi cấp/Đơn vị cấp: ........................................\n"
            "- Ngày cấp: ....................................................",
            27,
        ),
        (
            [
                "- Người đại diện/ủy quyền: {{sim_customer_name}}\n"
                "- Số ĐDCN/ĐDĐT/Hộ chiếu",
                2,
                ": {{sim_customer_id_number}}\n"
                "- Ngày tháng năm sinh: {{sim_customer_birth_date}}\n"
                "- Giới tính: {{sim_customer_gender}}\n"
                "- Ngày cấp: {{sim_customer_issue_date}}\n"
                "- Nơi cấp/Đơn vị cấp: {{sim_customer_issue_place}}\n"
                "- Địa chỉ theo giấy tờ dùng để đăng ký thông tin thuê bao: {{sim_customer_address}}\n"
                "- Quốc tịch: {{sim_customer_nationality}}",
            ],
            "- Người đại diện/ủy quyền: ....................................\n"
            "- Số ĐDCN/ĐDĐT/Hộ chiếu: ..................................\n"
            "- Ngày tháng năm sinh: ........................................\n"
            "- Giới tính: ........................................................\n"
            "- Ngày cấp: ........................................................\n"
            "- Nơi cấp/Đơn vị cấp: ...........................................\n"
            "- Địa chỉ theo giấy tờ dùng để đăng ký thông tin thuê bao: .................................................................\n"
            "- Quốc tịch: ☐ Việt Nam    ☐ Nước ngoài",
            49,
        ),
        (
            "- Nơi gửi thông báo cước và thanh toán: {{sim_customer_address}}",
            "- Nơi gửi thông báo cước và thanh toán: .................................................................",
            13,
        ),
        (
            "- Số điện thoại liên hệ: {{sim_customer_phone}}",
            "- Số điện thoại liên hệ: ........................................",
            10,
        ),
        (
            "- Email: {{sim_customer_email}}",
            "- Email: ........................................................",
            10,
        ),
    ]
    for left, right, height in rows:
        cells = table.add_row().cells
        _set_cell_width(cells[0], 95)
        _set_cell_width(cells[1], 95)
        if isinstance(left, list):
            _write_segments(cells[0], left, size=8.6)
        else:
            _write(cells[0], left, size=8.6)
        _write(cells[1], right, size=8.6)
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        table.rows[-1].height = Mm(height)
        table.rows[-1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        _set_cant_split(table.rows[-1])


def _service_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=0, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    widths = (11, 84, 11, 84)
    for left, right in rows:
        cells = table.add_row().cells
        for cell, width in zip(cells, widths):
            _set_cell_width(cell, width)
        _write(cells[0], "☐", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        _write(cells[1], left, size=8.7)
        _write(cells[2], "☐", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        _write(cells[3], right, size=8.7)
        table.rows[-1].height = Mm(10)
        table.rows[-1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        _set_cant_split(table.rows[-1])
    time_cells = table.add_row().cells
    merged = time_cells[0].merge(time_cells[3])
    _write(merged, "Thời gian thay đổi: Từ ngày: .................................................................", size=8.7)
    _set_cant_split(table.rows[-1])


def _build_page_two(document: Document) -> None:
    section = document.add_section(WD_SECTION.NEW_PAGE)
    _set_footer(section, [3])
    subscriber_line = _add_line(
        document,
        "Số thuê bao: ",
        "{{sim_subscriber_number}}        Sê-ri SIM hiện tại: ................................................",
        size=9,
        after=2,
    )
    _keep(subscriber_line, with_next=True)
    intro = document.add_paragraph()
    _format_paragraph(intro, after=2)
    _format_run(
        intro.add_run(
            "Yêu cầu thay đổi các dịch vụ như sau (Quý khách đánh dấu X vào ô vuông trước tên các dịch vụ cần "
            "đăng ký bổ sung, dấu O vào ô vuông trước tên các dịch vụ đề nghị cắt):"
        ),
        size=8.5,
    )
    _keep(intro, with_next=True)
    section_two = _section_heading(document, "II", "THAY ĐỔI DỊCH VỤ")
    _add_footnote_marker(section_two, 3)
    service_default = _add_line(document, "2.1. Các dịch vụ mặc định (Đánh dấu X vào ô dịch vụ tương ứng):", size=9, bold_label=True, after=1)
    _keep(service_default, with_next=True)
    _service_table(document, [
        ("Thoại, hiển thị số thuê bao chủ gọi, chuyển tiếp cuộc gọi, chờ cuộc gọi, cuộc gọi hội nghị", "GPRS"),
        ("Nhắn tin ngắn quốc tế", "Gọi khẩn cấp"),
        ("Nhắn tin ngắn trong nước", "Nhắn tin đa phương tiện"),
    ])
    service_optional = _add_line(document, "2.2. Các dịch vụ đăng ký (Đánh dấu X vào ô dịch vụ tương ứng):", size=9, bold_label=True, after=1)
    _keep(service_optional, with_next=True)
    _service_table(document, [
        ("In bảng kê chi tiết cước", "Chuyển vùng quốc tế"),
        ("Gọi quốc tế", "Hộp thư thoại"),
        ("Nhạc chuông chờ", "Dữ liệu linh hoạt"),
    ])

    sim_table = document.add_table(rows=2, cols=3)
    sim_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sim_table.autofit = False
    for row in sim_table.rows:
        for cell, width in zip(row.cells, (53, 92, 45)):
            _set_cell_width(cell, width)
            _set_cell_margins(cell, top=30, bottom=30)
        _set_cant_split(row)
    _write(sim_table.cell(0, 0), "III. THAY SIMCARD:", size=10, bold=True, color=ORANGE)
    _write(sim_table.cell(0, 1), "Số sê-ri SIM mới: {{sim_new_serial}}", size=9)
    _write(sim_table.cell(0, 2), "☐ Miễn phí", size=9)
    _write(sim_table.cell(1, 0), "Lý do:", size=9)
    _write(
        sim_table.cell(1, 1).merge(sim_table.cell(1, 2)),
        "{{sim_reason_lost_mark}} Mất SIM      {{sim_reason_damaged_mark}} Hỏng SIM      "
        "{{sim_reason_other_mark}} Lý do khác: {{sim_reason_other_value}}",
        size=9,
    )
    _add_line(document, "Thời gian yêu cầu thay đổi: Từ ", "{{sim_request_date_line}}", size=9, after=3)

    _add_line(document, "CAM KẾT CỦA KHÁCH HÀNG:", size=10, bold_label=True, after=0)
    _add_line(document, "Dưới đây là một số thông tin lịch sử quá trình sử dụng số thuê bao của tôi:", size=8.7, after=0)
    _add_line(
        document,
        "• Chi tiết 5 số thường xuyên liên lạc trong 3 tháng gần nhất: ",
        "1. {{frequent_phone_1}}; 2. {{frequent_phone_2}}; 3. {{frequent_phone_3}}; "
        "4. {{frequent_phone_4}}; 5. {{frequent_phone_5}}",
        size=8.4,
        after=0,
        indent_mm=5,
    )
    _add_line(document, "• Thời gian kích hoạt (sai số trong vòng 30 ngày): ", "{{activation_date}}", size=8.4, after=0, indent_mm=5)
    _add_line(
        document,
        "• Giá trị nạp thẻ gần nhất: ",
        "{{recent_topup_value}}    Hình thức nạp tiền: {{recent_topup_method}}",
        size=8.4,
        after=0,
        indent_mm=5,
    )
    _add_line(document, "• Thời hạn sử dụng còn lại: ", "{{remaining_validity}}", size=8.4, after=0, indent_mm=5)
    _add_line(document, "• Số tiền còn lại trong tài khoản: ", "{{account_balance}}", size=8.4, after=0, indent_mm=5)
    _add_line(document, "• Dịch vụ cuối cùng thay đổi: ", "{{last_changed_service}}", size=8.4, after=0, indent_mm=5)


def _signature_cell(cell, heading: str, *, given_key: str = "", full_key: str = "") -> None:
    """Every signature is 2 lines: "ký tên" (given_key -- just the given
    name, the actual signature stroke) above "ghi rõ họ tên" (full_key --
    the full name), both in the embedded Great Vibes script font. This
    replaces what used to be a real signature image (see font_embed.py)."""
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    _set_cell_margins(cell, top=100, bottom=100)
    p = cell.paragraphs[0]
    _format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    _format_run(p.add_run(heading), size=10.5, bold=True)
    _keep(p, with_next=True)
    p = cell.add_paragraph()
    _format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=3)
    _format_run(p.add_run("(Ký, ghi rõ họ tên)"), size=9, italic=True)
    _keep(p, with_next=bool(given_key or full_key))
    if given_key:
        given_p = cell.add_paragraph()
        _format_paragraph(given_p, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
        given_run = given_p.add_run("{{ " + given_key + " }}")
        _format_run(given_run, size=9.5, bold=True)
        apply_signature_font(given_run)
        _keep(given_p, with_next=bool(full_key))
    if full_key:
        full_p = cell.add_paragraph()
        _format_paragraph(full_p, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
        full_run = full_p.add_run("{{ " + full_key + " }}")
        _format_run(full_run, size=9.5, bold=True)
        apply_signature_font(full_run)


def _build_page_three(document: Document) -> None:
    section = document.add_section(WD_SECTION.NEW_PAGE)
    _set_footer(section, [])
    declaration = document.add_paragraph()
    _format_paragraph(declaration, after=5)
    _format_run(
        declaration.add_run(
            "Tôi xin cam đoan số thuê bao trên là của tôi sở hữu. Trường hợp có xảy ra khiếu kiện, tranh chấp "
            "quyền sở hữu hoặc có vấn đề gì liên quan đến số thuê bao trên, tôi xin chịu trách nhiệm trước pháp "
            "luật và Công ty Cổ phần Viễn thông Di động Vietnamobile có quyền thu hồi lại số thuê bao này theo "
            "quy định của pháp luật."
        ),
        size=9.5,
    )
    _keep(declaration, with_next=True)
    date = document.add_paragraph()
    _format_paragraph(date, align=WD_ALIGN_PARAGRAPH.RIGHT, after=1)
    _format_run(date.add_run("{{sim_document_date_line}}"), size=9)
    _keep(date, with_next=True)

    table = document.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for cell in row.cells:
            _set_cell_width(cell, 95)
        row.height = Mm(55)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        _set_cant_split(row)
    _signature_cell(
        table.cell(0, 0), "KHÁCH HÀNG ĐẠI DIỆN",
        given_key="customer_signature_given_name", full_key="customer_signature_name",
    )
    _signature_cell(
        table.cell(0, 1), "ĐẠI DIỆN BÊN CUNG CẤP DỊCH VỤ",
        given_key="provider_representative_given_name", full_key="provider_representative",
    )
    _signature_cell(
        table.cell(1, 0), "GIAO DỊCH VIÊN",
        given_key="sim_operator_signature_given_name", full_key="sim_operator_signature_name",
    )
    _signature_cell(table.cell(1, 1), "NHÂN VIÊN ĐẤU NỐI")


def build(output_docx: Path = OUTPUT_DOCX) -> Path:
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(17)
    section.bottom_margin = Mm(16)
    section.left_margin = Mm(10)
    section.right_margin = Mm(10)
    section.header_distance = Mm(2)
    section.footer_distance = Mm(2)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(9.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    embed_signature_font(document)
    _add_header(section)
    _set_footer(section, [1, 2])
    _build_page_one(document)
    _build_page_two(document)
    _build_page_three(document)

    document.core_properties.title = "Phiếu cung cấp và thay đổi dịch vụ trả trước"
    document.core_properties.subject = "Biểu mẫu Word native dựng lại từ PDF tháng 09/2025"
    document.core_properties.comments = "Không sử dụng ảnh nền hoặc ảnh chụp trang PDF."
    document.save(output_docx)
    return output_docx


if __name__ == "__main__":
    print(build())
