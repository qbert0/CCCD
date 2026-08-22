"""One-time, idempotent fix: style every printed signature-line name
(customer_signature_name, provider_representative, etc.) with the embedded
"Great Vibes" script font so it reads like a real signature instead of
plain Times New Roman -- see font_embed.py for why the font is embedded
into each docx rather than just referenced by name.

Targets are exact (table_index, row_index, col_index) cell coordinates,
found by dumping each template's tables and confirmed by hand -- NOT "every
occurrence of this placeholder name", since prepaid_contract's own
{{ provider_representative }} also appears once as plain prose ("Người đại
diện: {{ provider_representative }} Chức vụ: ...") that must stay in the
ordinary body font.

sim_change_form is excluded here on purpose: it's rebuilt from scratch by
build_sim_change_template.py, which already applies this styling itself,
so hand-patching its own docx here would just be overwritten on the next
regenerate.

Run from the repository root with the desktop virtual environment.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from desktop_app.backend.documents.font_embed import apply_signature_font, embed_signature_font  # noqa: E402

DOCS_DIR = ROOT / "desktop_app" / "backend" / "documents"

# path -> list of (table_index, row_index, col_index) signature-name cells.
TARGETS: dict[Path, list[tuple[int, int, int]]] = {
    DOCS_DIR / "transfer" / "00_MAU_BIEN_BAN_CHUYEN_CHU_QUYEN.docx": [(1, 2, 0), (1, 2, 1), (1, 2, 2)],
    DOCS_DIR / "aftersale" / "00_MAU_CAM_KET_SAU_BAN_HANG.docx": [(0, 2, 0), (0, 2, 1), (0, 2, 2)],
    DOCS_DIR / "beautiful_number" / "00_MAU_PHU_LUC_CAM_KET_SO_DEP_editable.docx": [(1, 2, 0), (1, 2, 1)],
    DOCS_DIR / "prepaid_contract" / "00_MAU_HOP_DONG_TRA_TRUOC.docx": [(1, 2, 0), (1, 2, 1)],
}


def apply_to(path: Path, cells: list[tuple[int, int, int]]) -> int:
    document = Document(path)
    embed_signature_font(document)
    styled = 0
    for table_index, row_index, col_index in cells:
        cell = document.tables[table_index].rows[row_index].cells[col_index]
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                apply_signature_font(run)
                styled += 1
    document.save(path)
    return styled


def main() -> None:
    for path, cells in TARGETS.items():
        count = apply_to(path, cells)
        print(f"{path.relative_to(ROOT)}: styled {count} run(s) in {len(cells)} cell(s)")


if __name__ == "__main__":
    main()
