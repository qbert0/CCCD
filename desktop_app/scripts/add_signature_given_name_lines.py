"""One-time, idempotent fix: give every signature its "ký tên" line (the
given name alone, in the embedded Great Vibes script font) in row 1 of its
signature table, right above the existing "ghi rõ họ tên" (full name) line
in row 2 -- replacing what used to be either a blank cell (parties who
never had a signature image) or a real signature image (the provider/
operator marker cells). See font_embed.py and renderer.py's
SIGNATURE_GIVEN_NAME_SOURCES for why.

sim_change_form is excluded here on purpose: it's rebuilt from scratch by
build_sim_change_template.py, which already does this itself.

Run from the repository root with the desktop virtual environment.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from desktop_app.backend.documents.font_embed import apply_signature_font, embed_signature_font  # noqa: E402

DOCS_DIR = ROOT / "desktop_app" / "backend" / "documents"

# path -> list of (table_index, row_index, col_index, given_name_key).
# row_index is always 1: row 0 is the heading/hint, row 2 the existing
# full-name line -- see the module docstring.
TARGETS: dict[Path, list[tuple[int, int, int, str]]] = {
    DOCS_DIR / "transfer" / "00_MAU_BIEN_BAN_CHUYEN_CHU_QUYEN.docx": [
        (1, 1, 0, "new_owner_signature_given_name"),
        (1, 1, 1, "customer_representative_signature_given_name"),
        (1, 1, 2, "provider_representative_given_name"),
    ],
    DOCS_DIR / "aftersale" / "00_MAU_CAM_KET_SAU_BAN_HANG.docx": [
        (0, 1, 0, "aftersale_requester_signature_given_name"),
        (0, 1, 1, "aftersale_new_owner_signature_given_name"),
        (0, 1, 2, "aftersale_clerk_signature_given_name"),
    ],
    DOCS_DIR / "beautiful_number" / "00_MAU_PHU_LUC_CAM_KET_SO_DEP_editable.docx": [
        (1, 1, 0, "customer_signature_given_name"),
        (1, 1, 1, "provider_representative_given_name"),
    ],
    DOCS_DIR / "prepaid_contract" / "00_MAU_HOP_DONG_TRA_TRUOC.docx": [
        (1, 1, 0, "prepaid_party_a_signature_given_name"),
        (1, 1, 1, "provider_representative_given_name"),
    ],
}


def apply_to(path: Path, cells: list[tuple[int, int, int, str]]) -> int:
    document = Document(path)
    embed_signature_font(document)
    filled = 0
    for table_index, row_index, col_index, given_key in cells:
        cell = document.tables[table_index].rows[row_index].cells[col_index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("{{ " + given_key + " }}")
        apply_signature_font(run)
        filled += 1
    document.save(path)
    return filled


def main() -> None:
    for path, cells in TARGETS.items():
        count = apply_to(path, cells)
        print(f"{path.relative_to(ROOT)}: filled {count} ký tên cell(s)")


if __name__ == "__main__":
    main()
