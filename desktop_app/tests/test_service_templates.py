import tempfile
import unittest
from pathlib import Path

from desktop_app.backend.documents import generate_service_template, required_input_numbers
from desktop_app.backend.domain.models import (
    SERVICE_TEMPLATE_DOCUMENTS,
    DocumentType,
    PersonData,
    ReportData,
    ServiceTemplate,
)


class ServiceTemplateTest(unittest.TestCase):
    def test_transfer_services_have_all_four_documents_and_sim_has_two(self):
        expected_transfer = [
            DocumentType.TRANSFER,
            DocumentType.AFTERSALE,
            DocumentType.BEAUTIFUL_NUMBER,
            DocumentType.PREPAID_CONTRACT,
        ]
        for template in ServiceTemplate:
            if template == ServiceTemplate.SIM_REPLACEMENT:
                self.assertEqual(
                    SERVICE_TEMPLATE_DOCUMENTS[template],
                    [DocumentType.AFTERSALE, DocumentType.SIM_CHANGE_FORM],
                )
            else:
                self.assertEqual(SERVICE_TEMPLATE_DOCUMENTS[template], expected_transfer)

    def test_numbered_image_requirements_follow_party_type(self):
        self.assertEqual(required_input_numbers(ServiceTemplate.PREPAID_TRANSFER_ORG), (1, 2, 3))
        self.assertEqual(required_input_numbers(ServiceTemplate.COMMITMENT_TRANSFER_ORG), (1, 2, 3))
        self.assertEqual(required_input_numbers(ServiceTemplate.SIM_REPLACEMENT), (1, 2, 3))
        self.assertEqual(
            required_input_numbers(ServiceTemplate.PREPAID_TRANSFER_INDIVIDUAL),
            (1, 2, 3, 4, 5, 6),
        )
        self.assertEqual(
            required_input_numbers(ServiceTemplate.COMMITMENT_TRANSFER_INDIVIDUAL),
            (1, 2, 3, 4, 5, 6),
        )

    def base_data(self, *, organization: bool) -> ReportData:
        customer = PersonData(
            full_name="NGUYỄN VĂN ĐẠI DIỆN",
            id_number="001099999999",
            date_of_birth="01/01/1990",
            nationality="Việt Nam",
            address="123 Xã Đàn, Hà Nội",
            issue_date="01/01/2021",
            issue_place="Cục Cảnh sát QLHC về TTXH",
        )
        if organization:
            customer.entity_type = "Tổ chức"
            customer.organization_name = "CÔNG TY TNHH ABC"
            customer.headquarters_address = "1 Xã Đàn, Hà Nội"
            customer.business_registration_number = "0101234567"
            customer.business_registration_issue_date = "01/01/2015"
            customer.business_registration_issue_place = "Sở KHĐT Hà Nội"
            customer.representative_name = "NGUYỄN VĂN ĐẠI DIỆN"
            customer.representative_position = "Giám đốc"
        return ReportData(
            document_type=DocumentType.TRANSFER,  # overwritten per document by generate_service_template
            customer=customer,
            new_owner=PersonData(
                full_name="TRẦN THỊ B",
                id_number="001204001289",
                date_of_birth="01/02/2004",
                nationality="Việt Nam",
                address="Hoàn Kiếm, Hà Nội",
                issue_date="02/02/2022",
                issue_place="Cục Cảnh sát QLHC về TTXH",
                phone="0909111222",
            ),
            subscribers=[
                {
                    "subscriber_number": "0925123456",
                    "monthly_fee": "500.000 đồng",
                    "activation_date": "19/08/2026",
                    "sim_serial": "8984041234567890001",
                    "commitment_note": "",
                }
            ],
            payment_method="Trả trước",
            transfer_effective_date="19/08/2026",
            backup_phone_1="0901234567",
            shop_id_number="0101234567",
            shop_issue_date="01/01/2020",
            shop_issue_place="Sở Tài chính Hà Nội",
            provider_representative="LÊ THỊ ĐẠI DIỆN",
            shop_address="123 Xã Đàn, Hà Nội",
            shop_phone="02412345678",
            service_point_name="Điểm giao dịch Xã Đàn",
            staff_name="Nguyễn Giao Dịch",
            registration_time="10:00",
            commitment_months="12",
        )

    def test_each_mau_generates_exactly_its_own_document_set(self):
        cases = [
            (ServiceTemplate.PREPAID_TRANSFER_ORG, True),
            (ServiceTemplate.PREPAID_TRANSFER_INDIVIDUAL, False),
            (ServiceTemplate.COMMITMENT_TRANSFER_INDIVIDUAL, False),
            (ServiceTemplate.COMMITMENT_TRANSFER_ORG, True),
            (ServiceTemplate.SIM_REPLACEMENT, False),
        ]
        for template, organization in cases:
            with self.subTest(template=template):
                data = self.base_data(organization=organization)
                if template == ServiceTemplate.SIM_REPLACEMENT:
                    data.service_action = "Thay SIM"
                with tempfile.TemporaryDirectory() as folder:
                    outputs, errors = generate_service_template(template, data, Path(folder))
                    self.assertEqual(errors, [])
                    self.assertEqual(len(outputs), len(SERVICE_TEMPLATE_DOCUMENTS[template]))
                    for output in outputs:
                        self.assertTrue(output.exists())

    def test_monthly_fee_required_only_for_commitment_mau(self):
        # "Cước cam kết tối thiểu/tháng" only makes sense as a hard
        # requirement for the 2 "cam kết" mẫu (3/4, khách cam kết mức cước
        # tối thiểu đổi lại số đẹp). The 2 "trả trước" mẫu (1/2) generate
        # the exact same appendix document but don't carry that commitment,
        # so the field must stay optional there.
        cases = [
            (ServiceTemplate.PREPAID_TRANSFER_INDIVIDUAL, False, True),
            (ServiceTemplate.PREPAID_TRANSFER_ORG, True, True),
            (ServiceTemplate.COMMITMENT_TRANSFER_INDIVIDUAL, False, False),
            (ServiceTemplate.COMMITMENT_TRANSFER_ORG, True, False),
        ]
        for template, organization, fee_optional in cases:
            with self.subTest(template=template):
                data = self.base_data(organization=organization)
                data.subscribers[0]["monthly_fee"] = ""
                with tempfile.TemporaryDirectory() as folder:
                    outputs, errors = generate_service_template(template, data, Path(folder))
                    fee_errors = [error for error in errors if "monthly_fee" in error.path]
                    if fee_optional:
                        self.assertEqual(fee_errors, [])
                        self.assertEqual(errors, [])
                    else:
                        self.assertTrue(fee_errors)
                        self.assertEqual(outputs, [])

    def test_missing_field_fails_the_whole_batch_atomically(self):
        data = self.base_data(organization=True)
        data.customer.business_registration_issue_date = ""  # required by TransferSchema
        with tempfile.TemporaryDirectory() as folder:
            outputs, errors = generate_service_template(
                ServiceTemplate.PREPAID_TRANSFER_ORG, data, Path(folder)
            )
            self.assertEqual(outputs, [])
            self.assertTrue(any(error.path == "customer.business_registration_issue_date" for error in errors))
            # Nothing partially written -- the whole point of validating first.
            self.assertEqual(list(Path(folder).iterdir()), [])

    def test_duplicate_subscriber_numbers_fail_the_whole_service_batch(self):
        data = self.base_data(organization=False)
        first = dict(data.subscribers[0])
        second = dict(first)
        second["subscriber_number"] = "0925 123 456"
        data.subscribers = [first, second]

        with tempfile.TemporaryDirectory() as folder:
            outputs, errors = generate_service_template(
                ServiceTemplate.PREPAID_TRANSFER_INDIVIDUAL, data, Path(folder)
            )
            duplicate_paths = {
                error.path for error in errors if "bị trùng" in error.message
            }
            self.assertEqual(outputs, [])
            self.assertIn("subscribers.0.subscriber_number", duplicate_paths)
            self.assertIn("subscribers.1.subscriber_number", duplicate_paths)
            self.assertEqual(list(Path(folder).iterdir()), [])

    def test_blank_row_activation_date_defaults_to_today_instead_of_blocking(self):
        # The minimal "just scan CCCD + type a subscriber number" flow (the
        # compact primary-number input, not the full subscriber table) can
        # produce a row with no activation_date at all -- it must default
        # to today rather than block Prepaid Contract generation.
        import datetime

        data = self.base_data(organization=False)
        data.subscribers = [{
            "subscriber_number": "0925123456", "monthly_fee": "500", "activation_date": "",
            "sim_serial": "", "commitment_note": "",
        }]
        with tempfile.TemporaryDirectory() as folder:
            outputs, errors = generate_service_template(
                ServiceTemplate.PREPAID_TRANSFER_INDIVIDUAL, data, Path(folder)
            )
            self.assertEqual(errors, [])
            self.assertEqual(len(outputs), len(SERVICE_TEMPLATE_DOCUMENTS[ServiceTemplate.PREPAID_TRANSFER_INDIVIDUAL]))
            today = datetime.date.today().strftime("%d/%m/%Y")
            from docx import Document
            document_types = SERVICE_TEMPLATE_DOCUMENTS[ServiceTemplate.PREPAID_TRANSFER_INDIVIDUAL]
            prepaid_output = outputs[document_types.index(DocumentType.PREPAID_CONTRACT)]
            table = Document(str(prepaid_output)).tables[0]
            self.assertIn(today, table.rows[1].cells[2].text)

    def test_prepaid_variant_gets_the_subscriber_list_populated(self):
        from docx import Document

        data = self.base_data(organization=False)
        with tempfile.TemporaryDirectory() as folder:
            outputs, errors = generate_service_template(
                ServiceTemplate.PREPAID_TRANSFER_INDIVIDUAL, data, Path(folder)
            )
            self.assertEqual(errors, [])
            document_types = SERVICE_TEMPLATE_DOCUMENTS[ServiceTemplate.PREPAID_TRANSFER_INDIVIDUAL]
            prepaid_output = outputs[document_types.index(DocumentType.PREPAID_CONTRACT)]
            table = Document(str(prepaid_output)).tables[0]
            self.assertIn("0925123456", table.rows[1].cells[0].text)

    def test_commitment_appendix_belongs_to_the_new_owner(self):
        from docx import Document

        data = self.base_data(organization=True)
        with tempfile.TemporaryDirectory() as folder:
            outputs, errors = generate_service_template(
                ServiceTemplate.COMMITMENT_TRANSFER_ORG, data, Path(folder)
            )
            self.assertEqual(errors, [])
            document_types = SERVICE_TEMPLATE_DOCUMENTS[ServiceTemplate.COMMITMENT_TRANSFER_ORG]
            appendix_output = outputs[document_types.index(DocumentType.BEAUTIFUL_NUMBER)]
            document = Document(str(appendix_output))
            content = "\n".join(
                [paragraph.text for paragraph in document.paragraphs]
                + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            )
            self.assertIn("Trần Thị B", content)
            self.assertIn("001204001289", content)

    def test_sim_replacement_creates_commitment_and_native_word_change_form(self):
        from docx import Document
        from docx.oxml.ns import qn
        from zipfile import ZipFile

        data = self.base_data(organization=False)
        data.service_action = "Thay SIM"
        with tempfile.TemporaryDirectory() as folder:
            outputs, errors = generate_service_template(
                ServiceTemplate.SIM_REPLACEMENT, data, Path(folder)
            )
            self.assertEqual(errors, [])
            self.assertEqual(len(outputs), 2)
            change_form = outputs[1]
            document = Document(change_form)
            # The form is native Word content: five real tables, and each
            # of its 3 pages is its own section (needed so each page's
            # footer can carry its own footnote text -- see below). It
            # must never regress to three pasted screenshots pretending to
            # be a DOCX.
            self.assertEqual(len(document.tables), 5)
            self.assertEqual(len(document.sections), 3)
            # Email belongs to the customer-information block instead of a
            # detached table at the top of page two.
            self.assertIn("Email", document.tables[0].rows[-1].cells[0].text)
            # Keep every structured row together. This prevents the bottom
            # half of a customer/signature row from spilling onto a new page.
            self.assertEqual(
                len(document.element.findall(".//" + qn("w:trPr") + "/" + qn("w:cantSplit"))),
                20,
            )
            with ZipFile(change_form) as archive:
                names = archive.namelist()
                images = [
                    name for name in archive.namelist()
                    if name.startswith("word/media/image") and name.endswith(".png")
                ]
                document_xml = archive.read("word/document.xml")
                relationships_xml = archive.read("word/_rels/document.xml.rels")
            # Footnote text is printed as real per-page FOOTER content
            # (pinned to the page bottom, above the brand banner), never
            # real OOXML w:footnoteReference/footnotes.xml: LibreOffice's
            # headless PDF export (this app's own render pipeline -- see
            # docx_to_images.py) was found to silently drop or misattach a
            # footnote's body text once more than one crosses a hard page
            # break, so real footnotes can never safely be used here.
            self.assertNotIn("word/footnotes.xml", names)
            self.assertEqual(document_xml.count(b"<w:footnoteReference"), 0)
            self.assertNotIn(b"relationships/footnotes", relationships_xml)
            footer_text_by_section = [
                "\n".join(p.text for p in section.footer.paragraphs)
                for section in document.sections
            ]
            self.assertIn("Số Quyết định thành lập", footer_text_by_section[0])
            self.assertIn("Số Định danh cá nhân", footer_text_by_section[0])
            self.assertIn("Các nội dung bỏ trống tại Phần II, III", footer_text_by_section[1])
            # Page 3 has no footnote of its own -- its footer must not
            # carry over page 1 or 2's text (each section's footer is
            # explicitly unlinked from the previous one for exactly this
            # reason).
            for phrase in (
                "Số Quyết định thành lập",
                "Số Định danh cá nhân",
                "Các nội dung bỏ trống tại Phần II, III",
            ):
                self.assertNotIn(phrase, footer_text_by_section[2])
            self.assertNotIn("Các nội dung bỏ trống tại Phần II, III", footer_text_by_section[0])
            # Exactly the header/footer brand banners now (reused from the
            # transfer template's own real graphics, not a reconstruction)
            # -- every signature is text (the embedded Great Vibes script
            # font) rather than a real signature image now, so there's no
            # third image. No full-page screenshots pretending to be the
            # document either way.
            self.assertEqual(len(images), 2)

    def test_sim_change_form_customer_phone_is_the_shop_phone(self):
        # "Số điện thoại liên hệ" on this form is the SHOP's own contact
        # number (auto-defaulted from the company profile via
        # data.shop_phone), not the customer's -- confirmed directly by the
        # user, since a clerk shouldn't have to type it per case.
        from docx import Document

        data = self.base_data(organization=False)
        data.service_action = "Thay SIM"
        data.shop_phone = "02412345678"
        with tempfile.TemporaryDirectory() as folder:
            outputs, errors = generate_service_template(
                ServiceTemplate.SIM_REPLACEMENT, data, Path(folder)
            )
            self.assertEqual(errors, [])
            document_types = SERVICE_TEMPLATE_DOCUMENTS[ServiceTemplate.SIM_REPLACEMENT]
            change_form = outputs[document_types.index(DocumentType.SIM_CHANGE_FORM)]
            document = Document(str(change_form))
            content = "\n".join(
                [paragraph.text for paragraph in document.paragraphs]
                + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            )
            self.assertIn("Số điện thoại liên hệ: 02412345678", content)

    def test_sim_replacement_rejects_more_than_one_subscriber(self):
        # The shared subscriber-list editor lets every service add rows via
        # "+ Thêm số thuê bao khác" (needed by the other 4 services), but
        # Thay SIM's physical form has exactly one "Số thuê bao" line -- a
        # 2nd row would silently be dropped rather than actually block
        # generation, hiding a real data-entry mistake.
        data = self.base_data(organization=False)
        data.service_action = "Thay SIM"
        data.subscribers = [
            dict(data.subscribers[0]),
            {**data.subscribers[0], "subscriber_number": "0987654321"},
        ]
        with tempfile.TemporaryDirectory() as folder:
            outputs, errors = generate_service_template(
                ServiceTemplate.SIM_REPLACEMENT, data, Path(folder)
            )
            self.assertEqual(outputs, [])
            self.assertTrue(any("một thuê bao" in error.message for error in errors))

    def test_document_types_override_narrows_the_generated_set(self):
        # The settings-driven document-set override (WebBridge's
        # get/save_document_set_settings) resolves to an explicit list
        # before calling here -- this is the seam it plugs into.
        data = self.base_data(organization=False)
        with tempfile.TemporaryDirectory() as folder:
            outputs, errors = generate_service_template(
                ServiceTemplate.PREPAID_TRANSFER_INDIVIDUAL, data, Path(folder),
                document_types=[DocumentType.TRANSFER, DocumentType.AFTERSALE],
            )
            self.assertEqual(errors, [])
            self.assertEqual(len(outputs), 2)
            self.assertEqual({output.parent for output in outputs}, {Path(folder)})


if __name__ == "__main__":
    unittest.main()
