import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

from desktop_app.backend.documents import DocumentRegistry
from desktop_app.backend.documents.renderer import CHECKED_BOX, EMPTY_BOX, _PLACEHOLDER, _iter_paragraphs
from desktop_app.backend.domain.models import DocumentType, PersonData, ReportData


def _textbox_text(paragraph: Paragraph) -> list[str]:
    texts = []
    for txbx_content in paragraph._p.findall(".//" + qn("w:txbxContent")):
        for p_element in txbx_content.findall(qn("w:p")):
            texts.append(Paragraph(p_element, paragraph._parent).text)
    return texts


def docx_text(path: Path) -> str:
    document = Document(str(path))
    values = []
    for paragraph in document.paragraphs:
        values.append(paragraph.text)
        values.extend(_textbox_text(paragraph))
    for table in document.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)
    return "\n".join(values)


class DocumentTest(unittest.TestCase):
    def setUp(self):
        self.customer = PersonData(
            full_name="NGUYỄN VĂN AN",
            id_number="001099999999",
            date_of_birth="01/01/1999",
            nationality="Việt Nam",
            address="123 Đường Láng, Đống Đa, Hà Nội",
            issue_date="01/01/2021",
            issue_place="Cục Cảnh sát QLHC về TTXH",
            phone="0901234567",
        )
        self.registry = DocumentRegistry()

    def report(self, doc_type: DocumentType) -> ReportData:
        return ReportData(
            document_type=doc_type,
            customer=self.customer,
            new_owner=PersonData(
                full_name="TRẦN THỊ B",
                id_number="001204001289",
                date_of_birth="01/02/2004",
                nationality="Việt Nam",
                address="Hoàn Kiếm, Hà Nội",
                issue_date="02/02/2022",
                issue_place="Cục Cảnh sát QLHC về TTXH",
            ),
            document_date="11/08/2026",
            subscriber_number="0925123456",
            sim_serial="8984041234567890123",
            activation_date="11/08/2026",
            source_contract_number="HĐ-2026-01",
            source_contract_date="10/08/2026",
            transfer_effective_date="11/08/2026",
            provider_representative="LÊ THỊ ĐẠI DIỆN",
            shop_name="Vietnamobile",
            shop_address="123 Xã Đàn, Hà Nội",
            shop_phone="02435730123",
            shop_id_number="0101234567",
            shop_issue_date="01/01/2020",
            shop_issue_place="Sở Tài chính Hà Nội",
            staff_name="Nguyễn Giao Dịch",
            backup_phone_1="0901234567",
            service_point_name="Điểm giao dịch Xã Đàn",
            registration_time="09:30 ngày 11/08/2026",
            commitment_months="12 tháng",
            monthly_fee="500.000 đồng",
        )

    def test_generates_beautiful_number_docx_template(self):
        # Beautiful Number moved from a PDF/raster overlay to a real docx
        # template -- native text/tables, not an image, so a printed row is
        # never constrained by scanned-page geometry the way the old
        # template was.
        with tempfile.TemporaryDirectory() as folder:
            data = self.report(DocumentType.BEAUTIFUL_NUMBER)
            output = self.registry.for_data(data).generate(data, Path(folder))
            self.assertEqual(output.suffix, ".docx")
            text = docx_text(output)
            self.assertIn("0925123456", text)
            self.assertIn("Nguyễn Văn An", text)
            self.assertIn("001099999999", text)
            self.assertIn("500.000đ", text)
            self.assertNotIn("{{", text)

    def test_beautiful_number_legacy_row2_fields_produce_a_real_second_table_row(self):
        with tempfile.TemporaryDirectory() as folder:
            data = self.report(DocumentType.BEAUTIFUL_NUMBER)
            data.subscriber_number_1 = data.subscriber_number
            data.subscriber_number_2 = "0987654321"
            data.commitment_months_2 = "24"
            data.monthly_fee_2 = "300.000 đồng"
            output = self.registry.for_data(data).generate(data, Path(folder))
            document = Document(str(output))
            table = document.tables[0]
            self.assertEqual(len(table.rows), 3)  # header + 2 real data rows
            self.assertIn("0925123456", table.rows[1].cells[1].text)
            self.assertIn("12 tháng", table.rows[1].cells[2].text)
            self.assertIn("0987654321", table.rows[2].cells[1].text)
            self.assertIn("24 tháng", table.rows[2].cells[2].text)
            self.assertIn("300.000đ", table.rows[2].cells[3].text)

    def test_beautiful_number_many_rows_all_clone_into_the_real_table(self):
        with tempfile.TemporaryDirectory() as folder:
            data = self.report(DocumentType.BEAUTIFUL_NUMBER)
            data.beautiful_subscribers = [
                {
                    "subscriber_number": f"09251234{index:02d}",
                    "commitment_months": str(12 + index),
                    "monthly_fee": str(500 + index),
                    "commitment_note": f"Dòng {index + 1}",
                }
                for index in range(6)
            ]
            output = self.registry.for_data(data).generate(data, Path(folder))
            document = Document(str(output))
            table = document.tables[0]
            self.assertEqual(len(table.rows), 7)  # header + 6 cloned data rows
            self.assertIn("0925123400", table.rows[1].cells[1].text)
            self.assertIn("0925123405", table.rows[6].cells[1].text)
            self.assertIn("500.000đ", table.rows[1].cells[3].text)
            self.assertIn("505.000đ", table.rows[6].cells[3].text)
            self.assertEqual(table.rows[6].cells[0].text, "6")

    def test_beautiful_number_second_row_is_optional_until_started(self):
        data = self.report(DocumentType.BEAUTIFUL_NUMBER)
        errors = self.registry.for_data(data).check(data)
        self.assertFalse(any(error.path.endswith("_2") for error in errors))

        data.subscriber_number_2 = "0987654321"
        errors = self.registry.for_data(data).check(data)
        self.assertTrue(any(error.path == "commitment_months_2" for error in errors))

    def test_beautiful_number_months_accept_digits_without_a_unit(self):
        data = self.report(DocumentType.BEAUTIFUL_NUMBER)
        data.commitment_months = "mười hai"
        errors = self.registry.for_data(data).check(data)
        self.assertTrue(any(error.path == "commitment_months" for error in errors))

        data.commitment_months = "12"
        errors = self.registry.for_data(data).check(data)
        self.assertFalse(any(error.path == "commitment_months" for error in errors))

    def test_beautiful_number_fee_accepts_positive_thousands_only(self):
        data = self.report(DocumentType.BEAUTIFUL_NUMBER)
        data.monthly_fee = "năm trăm"
        errors = self.registry.for_data(data).check(data)
        self.assertTrue(any(error.path == "monthly_fee" for error in errors))

        data.monthly_fee = "500"
        errors = self.registry.for_data(data).check(data)
        self.assertFalse(any(error.path == "monthly_fee" for error in errors))

    def test_prepaid_individual_does_not_fill_organization_section(self):
        with tempfile.TemporaryDirectory() as folder:
            data = self.report(DocumentType.PREPAID_CONTRACT)
            output = self.registry.for_data(data).generate(data, Path(folder))
            text = docx_text(output)
            self.assertEqual(text.count("001099999999"), 1)
            self.assertIn("NGUYỄN VĂN AN", text)
            self.assertIn("8984041234567890123", text)
            self.assertIn("Điểm giao dịch Xã Đàn", text)
            self.assertNotIn("{{", text)

    def test_prepaid_template_accepts_an_intentionally_removed_placeholder(self):
        with tempfile.TemporaryDirectory() as folder:
            data = self.report(DocumentType.PREPAID_CONTRACT)
            module = self.registry.for_data(data)
            document = Document(str(module.template))
            removed = False
            for paragraph in _iter_paragraphs(document):
                updated = _PLACEHOLDER.sub(
                    lambda match: "" if match.group(1) == "prepaid_individual_email" else match.group(0),
                    paragraph.text,
                )
                if updated != paragraph.text:
                    paragraph.text = updated
                    removed = True
            self.assertTrue(removed)

            template = Path(folder) / "prepaid_without_email_placeholder.docx"
            document.save(str(template))
            module.template = template
            output = module.generate(data, Path(folder) / "output")
            self.assertNotIn("{{", docx_text(output))

    def test_empty_fields_use_dots_outside_tables_and_stay_blank_inside(self):
        with tempfile.TemporaryDirectory() as folder:
            transfer = self.report(DocumentType.TRANSFER)
            transfer.customer.authorization_number = ""
            transfer.customer.authorization_date = ""
            output = self.registry.for_data(transfer).generate(transfer, Path(folder))
            table = Document(str(output)).tables[0]
            self.assertEqual(table.rows[5].cells[1].text, "")
            self.assertEqual(table.rows[5].cells[2].text, "")

        with tempfile.TemporaryDirectory() as folder:
            aftersale = self.report(DocumentType.AFTERSALE)
            aftersale.other_attachment = ""
            output = self.registry.for_data(aftersale).generate(aftersale, Path(folder))
            self.assertIn("." * 48, docx_text(output))

        with tempfile.TemporaryDirectory() as folder:
            beautiful = self.report(DocumentType.BEAUTIFUL_NUMBER)
            beautiful.commitment_note = ""
            output = self.registry.for_data(beautiful).generate(beautiful, Path(folder))
            self.assertEqual(Document(str(output)).tables[0].rows[1].cells[4].text, "")

        with tempfile.TemporaryDirectory() as folder:
            prepaid = self.report(DocumentType.PREPAID_CONTRACT)
            prepaid.customer.email = ""
            output = self.registry.for_data(prepaid).generate(prepaid, Path(folder))
            self.assertIn("Email: " + "." * 24, docx_text(output))
            subscriber_table = Document(str(output)).tables[0]
            self.assertTrue(all(cell.text == "" for cell in subscriber_table.rows[2].cells))

    def test_generates_transfer_docx_with_customer(self):
        with tempfile.TemporaryDirectory() as folder:
            data = self.report(DocumentType.TRANSFER)
            output = self.registry.for_data(data).generate(data, Path(folder))
            text = docx_text(output)
            self.assertIn("NGUYỄN VĂN AN", text)
            self.assertIn("001099999999", text)
            self.assertIn("TRẦN THỊ B", text)
            self.assertIn("001204001289", text)
            self.assertIn("0925123456", text)
            # The baked dotted markers around placeholder values (see
            # scripts/bake_field_highlighting.py) were removed per direct
            # instruction (scripts/remove_baked_dots.py) -- a filled
            # value now reads directly, no surrounding "....." anymore.
            self.assertIn("HĐ-2026-01 ngày 10 tháng 08 năm 2026", text)
            self.assertIn("Hôm nay, ngày 11 tháng 08 năm 2026", text)
            self.assertNotIn("{{", text)
            self.assertNotIn("{{pay_method}}", text)
            self.assertNotIn("Cty TNHH Viễn Thông Tiến Prô", text)
            self.assertNotIn("Trụ sở chính\nBa Đình, Bỉm Sơn", text)

    def test_transfer_with_multiple_subscribers_clones_a_real_row_per_number(self):
        with tempfile.TemporaryDirectory() as folder:
            data = self.report(DocumentType.TRANSFER)
            data.subscribers = [
                {"subscriber_number": "0925123456"},
                {"subscriber_number": "0925111222"},
                {"subscriber_number": "0925333444"},
            ]
            output = self.registry.for_data(data).generate(data, Path(folder))
            table = Document(str(output)).tables[0]
            data_rows = table.rows[-3:]
            self.assertEqual([row.cells[0].text for row in data_rows], [
                "Số thuê bao 1", "Số thuê bao 2", "Số thuê bao 3",
            ])
            for row, number in zip(data_rows, ("0925123456", "0925111222", "0925333444")):
                self.assertEqual(row.cells[1].text, number)
                self.assertEqual(row.cells[2].text, number)
            text = docx_text(output)
            self.assertIn("0925123456, 0925111222, 0925333444", text)

    def test_aftersale_with_multiple_subscribers_joins_them_into_the_existing_blank(self):
        with tempfile.TemporaryDirectory() as folder:
            data = self.report(DocumentType.AFTERSALE)
            data.service_action = "Chuyển chủ quyền"
            data.subscribers = [
                {"subscriber_number": "0925123456"},
                {"subscriber_number": "0925111222"},
            ]
            output = self.registry.for_data(data).generate(data, Path(folder))
            text = docx_text(output)
            self.assertIn("0925123456, 0925111222", text)
            # Subscriber numbers still join into the existing sentence
            # blank, not a table -- the one table in this document now is
            # the (unrelated) signature block.
            self.assertEqual(len(Document(str(output)).tables), 1)

    def test_aftersale_identity_block_uses_the_actual_organization_requester(self):
        # The form labels this block "Khách hàng/Người yêu cầu". An
        # organization transfer therefore prints the old owner's company,
        # not an unrelated global shop identity.
        with tempfile.TemporaryDirectory() as folder:
            data = self.report(DocumentType.AFTERSALE)
            data.customer = PersonData(
                entity_type="Tổ chức",
                organization_name="CÔNG TY CHỦ CŨ ABC",
                business_registration_number="9999999999",
                business_registration_issue_date="01/01/2019",
                business_registration_issue_place="Sở KHĐT nào đó",
                headquarters_address="Địa chỉ chủ cũ",
            )
            data.service_action = "Chuyển chủ quyền"
            output = self.registry.for_data(data).generate(data, Path(folder))
            text = docx_text(output)
            self.assertIn("CÔNG TY CHỦ CŨ ABC", text)
            self.assertIn("9999999999", text)
            self.assertIn("Địa chỉ chủ cũ", text)
            self.assertNotIn("0101234567", text)

    def test_transfer_and_prepaid_print_the_actual_provider_representative(self):
        # Bên B's signature name used to be a literal "VÕ DUY NHẬT" baked
        # into the template regardless of who was actually signing -- both
        # templates now carry a real {{ provider_representative }} token
        # (plus its {{ provider_representative_given_name }} "ký tên" line,
        # replacing what used to be the company stamp+signature image), so
        # the printed name must track data.provider_representative and the
        # hardcoded name must never appear.
        with tempfile.TemporaryDirectory() as folder:
            transfer = self.report(DocumentType.TRANSFER)
            output = self.registry.for_data(transfer).generate(transfer, Path(folder))
            text = docx_text(output)
            # Bên B's printed full name is plain bold Times New Roman now
            # (per direct instruction), not the title-cased Great Vibes
            # script every other signature line uses -- so it stays
            # all-caps like the rest of the document's plain text.
            self.assertIn("LÊ THỊ ĐẠI DIỆN", text)
            self.assertIn("Diện", text)  # the "ký tên" given-name line
            self.assertNotIn("VÕ DUY NHẬT", text)

        with tempfile.TemporaryDirectory() as folder:
            prepaid = self.report(DocumentType.PREPAID_CONTRACT)
            output = self.registry.for_data(prepaid).generate(prepaid, Path(folder))
            text = docx_text(output)
            self.assertIn("LÊ THỊ ĐẠI DIỆN", text)
            self.assertIn("Diện", text)
            self.assertNotIn("VÕ DUY NHẬT", text)

    def test_transfer_template_signature_table_prints_each_party_and_hour_only(self):
        # The signature block used to be three absolutely-positioned
        # floating text boxes -- fragile to edit and impossible to add a
        # signature image into without breaking the layout (see the fixed
        # prepaid_contract equivalent). It's a real table now: row 0 =
        # heading/hint, row 1 = ký tên + họ tên stacked as two paragraphs
        # in the same cell (this table gets hand-edited in Word/LibreOffice
        # from time to time -- it started as 3 rows and settled at 2, still
        # keeping every party's own 2 lines together). Column ORDER isn't
        # asserted either -- the shop is free to reorder its own columns --
        # only that each of the 3 "Đại diện Bên X" columns still carries
        # the right content, wherever it physically sits.
        with tempfile.TemporaryDirectory() as folder:
            data = self.report(DocumentType.TRANSFER)
            data.customer.full_name = "NGUYỄN ĐẠI DIỆN BÊN A"
            data.transfer_time = "03:45"  # old saved state is normalized while rendering
            module = self.registry.for_data(data)
            output = module.preview(data, Path(folder))
            document = Document(str(output))
            text = docx_text(output)
            self.assertIn("từ 03 giờ", text)
            self.assertNotIn("03:45", text)
            self.assertIn("NGUYỄN ĐẠI DIỆN BÊN A", text)

            signature_table = next(
                t for t in document.tables
                if any("Đại diện Bên A" in cell.text.splitlines()[0] for cell in t.rows[0].cells)
            )
            self.assertEqual(len(signature_table.rows), 2)
            column_count = len(signature_table.rows[0].cells)
            cells_by_party = {
                signature_table.rows[0].cells[col].text.splitlines()[0]: "\n".join(
                    signature_table.rows[row].cells[col].text for row in range(len(signature_table.rows))
                )
                for col in range(column_count)
            }
            self.assertEqual(set(cells_by_party), {"Đại diện Bên A", "Đại diện Bên B", "Đại diện Bên C"})
            # Bên A/C's signature-line names print in the embedded Great
            # Vibes script font, which is title-cased rather than the
            # all-caps used everywhere else (see
            # renderer.py::_replace_placeholders) -- a full-caps run
            # through a cursive connecting face reads as tangled,
            # illegible strokes instead of a real signature. Bên B is
            # plain bold Times New Roman instead (per direct instruction),
            # so it keeps the all-caps the rest of the document uses.
            self.assertIn("Nguyễn Đại Diện Bên A", cells_by_party["Đại diện Bên A"])
            self.assertIn("LÊ THỊ ĐẠI DIỆN", cells_by_party["Đại diện Bên B"])
            # Bên C is the new owner's own name (new_owner_signature_name),
            # a variable independent from Bên B's provider_representative --
            # an earlier round briefly had Bên C reuse Bên B's exact
            # variable, which made the two cells always print the identical
            # value (a real bug the user caught: "nó có giá trị giống nhau
            # này"). They must stay different fields even when, by
            # coincidence, they'd render the same text.
            self.assertIn("Trần Thị B", cells_by_party["Đại diện Bên C"])
            self.assertNotIn("LÊ THỊ ĐẠI DIỆN", cells_by_party["Đại diện Bên C"])

    def test_prepaid_contract_signature_table_bien_a_is_not_bien_b(self):
        # Bên A and Bên B used to share one token ({{ provider_representative
        # }} in both cells) -- the same class of bug already fixed for
        # transfer's Bên C. Bên A must print whoever is actually taking over
        # the subscription (the customer signing for themselves in legacy
        # mode), independent of Bên B's fixed provider identity.
        with tempfile.TemporaryDirectory() as folder:
            data = self.report(DocumentType.PREPAID_CONTRACT)
            output = self.registry.for_data(data).generate(data, Path(folder))
            document = Document(str(output))
            signature_table = next(
                t for t in document.tables
                if any(
                    cell.text and "ĐẠI DIỆN BÊN A" in cell.text.splitlines()[0]
                    for cell in t.rows[0].cells
                )
            )
            self.assertEqual(len(signature_table.rows), 2)
            cells_by_party = {
                signature_table.rows[0].cells[col].text.splitlines()[0]: "\n".join(
                    signature_table.rows[row].cells[col].text for row in range(len(signature_table.rows))
                )
                for col in range(len(signature_table.rows[0].cells))
            }
            # Bên A's name prints in the embedded Great Vibes script font,
            # which is title-cased rather than the all-caps used elsewhere
            # (see renderer.py::_replace_placeholders). Bên B is plain bold
            # Times New Roman instead (per direct instruction), so it stays
            # all-caps like the rest of the document.
            self.assertIn("Nguyễn Văn An", cells_by_party["ĐẠI DIỆN BÊN A"])
            self.assertIn("LÊ THỊ ĐẠI DIỆN", cells_by_party["ĐẠI DIỆN BÊN B"])
            self.assertNotIn("LÊ THỊ ĐẠI DIỆN", cells_by_party["ĐẠI DIỆN BÊN A"])

            # The signature row must hold a fixed minimum height even when a
            # column carries no image (Bên A never does) -- otherwise there's
            # no blank space left for a physical signature.
            signature_row_tr_pr = signature_table.rows[1]._tr.find(qn("w:trPr"))
            row_height = signature_row_tr_pr.find(qn("w:trHeight"))
            self.assertEqual(row_height.get(qn("w:hRule")), "exact")
            self.assertGreater(int(row_height.get(qn("w:val"))), 1000)

    def test_transfer_rejects_same_current_and_new_owner(self):
        data = self.report(DocumentType.TRANSFER)
        data.new_owner.id_number = data.customer.id_number
        errors = self.registry.for_data(data).check(data)
        self.assertTrue(any(error.path == "new_owner.id_number" for error in errors))

    def test_transfer_contract_and_registration_form_fields_are_optional(self):
        data = self.report(DocumentType.TRANSFER)
        data.source_contract_number = ""
        data.source_contract_date = ""
        data.registration_form_date = ""
        errors = self.registry.for_data(data).check(data)
        optional_paths = {
            "source_contract_number", "source_contract_date", "registration_form_date",
        }
        self.assertFalse(any(error.path in optional_paths for error in errors))

    def test_prepaid_contact_service_point_and_registration_time_are_optional(self):
        data = self.report(DocumentType.PREPAID_CONTRACT)
        data.prepaid_structured_parties = True
        optional_paths = {
            "representative.phone", "new_owner.phone",
            "service_point_name", "registration_time",
        }
        required = set(self.registry.for_data(data).required_paths(data))
        self.assertTrue(optional_paths.isdisjoint(required))

    def test_transfer_rejects_minutes_or_an_hour_outside_the_day(self):
        data = self.report(DocumentType.TRANSFER)
        for invalid in ("03:45", "24"):
            data.transfer_time = invalid
            errors = self.registry.for_data(data).check(data)
            self.assertTrue(any(error.path == "transfer_time" for error in errors))

    def test_generates_aftersale_with_reviewed_fields(self):
        with tempfile.TemporaryDirectory() as folder:
            data = self.report(DocumentType.AFTERSALE)
            data.document_date = "11/08/2026"
            data.shop_address = "123 Xã Đàn, Hà Nội"
            data.shop_phone = "02435730123"
            output = self.registry.for_data(data).generate(data, Path(folder))
            text = docx_text(output)
            self.assertIn("Ngày 11 tháng 08 năm 2026", text)
            self.assertIn("Shop Xã Đàn", text)
            # The template was hand-edited to drop {{ aftersale_shop_address
            # }}/{{ aftersale_shop_phone }} entirely -- the shop no longer
            # wants this auto-filled here, just a blank line to write on by
            # hand. data.shop_address/shop_phone (set above) simply go
            # unused for this document type now; confirm the blank stays a
            # plain dotted line, not a leftover unresolved token.
            self.assertNotIn("{{ aftersale_shop_address }}", text)
            self.assertIn("0925123456", text)
            # Individual services and SIM replacement must print the
            # scanned person, while organization services print the actual
            # organization requester.
            self.assertIn("Số CMND/CCCD: 001099999999", text)
            self.assertNotIn("Người đại diện:", text)
            self.assertIn("NGUYỄN VĂN AN", text)
            self.assertIn("Nguyễn Giao Dịch", text)
            self.assertNotIn("{{", text)
            self.assertNotIn("CÔNG TY TNHH VIỄN THÔNG TIẾN PRÔ", text)

    def test_aftersale_clerk_signature_has_both_ky_ten_and_full_name_lines(self):
        # Every signature is 2 lines now instead of a real signature image:
        # "ký tên" (just the given name, the actual signature stroke) above
        # "ghi rõ họ tên" (the full name) -- both in the embedded Great
        # Vibes script font.
        with tempfile.TemporaryDirectory() as folder:
            data = self.report(DocumentType.AFTERSALE)
            output = self.registry.for_data(data).generate(data, Path(folder))
            document = Document(str(output))
            self.assertIn("Nguyễn Giao Dịch", docx_text(output))
            given_name_runs = [
                run for paragraph in document.paragraphs + [
                    p for table in document.tables for row in table.rows
                    for cell in row.cells for p in cell.paragraphs
                ]
                for run in paragraph.runs
                if run.text == "Dịch"
            ]
            self.assertTrue(given_name_runs)
            self.assertEqual(given_name_runs[0].font.name, "Great Vibes")

    def test_aftersale_only_fills_the_selected_service(self):
        with tempfile.TemporaryDirectory() as folder:
            data = self.report(DocumentType.AFTERSALE)
            data.service_action = "Chuyển chủ quyền"
            data.shop_phone_2 = "0987000111"
            data.backup_phone_2 = "0900000000"  # legacy field must no longer win
            output = self.registry.for_data(data).generate(data, Path(folder))
            text = docx_text(output)
            self.assertIn("☑ Chuyển chủ quyền", text)
            self.assertIn("☐ Cập nhật thông tin", text)
            self.assertIn("☐ Thay SIM", text)
            self.assertRegex(text, r"Vietnamobile: \.{20,}")
            # Filled values used to carry baked dotted markers (see
            # scripts/bake_field_highlighting.py) but those were removed
            # per direct instruction (scripts/remove_baked_dots.py) -- a
            # filled value now reads directly, no surrounding ".....".
            self.assertIn("cho Ông/Bà TRẦN THỊ B", text)
            self.assertIn("số điện thoại 2: 0987000111", text)
            self.assertNotIn("số điện thoại 2: 0900000000", text)

            document = Document(output)
            # DejaVu Sans (bundled with LibreOffice) draws U+2611 as a hollow
            # outline box with a small check actually inside it -- verified
            # by rendering a real generated document -- unlike most fonts,
            # which fall back to a colored, filled-looking emoji glyph.
            checkbox_runs = [
                run
                for paragraph in document.paragraphs
                for run in paragraph.runs
                if run.text in {CHECKED_BOX, EMPTY_BOX}
            ]
            self.assertTrue(checkbox_runs)
            self.assertTrue(all(run.font.name == "DejaVu Sans" for run in checkbox_runs))
            self.assertTrue(all(run.font.size and run.font.size.pt == 15 for run in checkbox_runs))
            for run in checkbox_runs:
                properties = run._r.rPr
                self.assertIsNone(properties.find(qn("w:w")))
                self.assertIsNone(properties.find(qn("w:spacing")))
                self.assertEqual(properties.find(qn("w:sz")).get(qn("w:val")), "30")
                self.assertEqual(properties.find(qn("w:szCs")).get(qn("w:val")), "30")
            # The signature block is a real table now (was a 3-column
            # SECTION with linear paragraph flow -- the same class of
            # fragility already fixed for the other 4 templates' signature
            # blocks). Row 0 = heading/hint, row 1 = ký tên + họ tên
            # stacked as two paragraphs in the same cell, so the 3 names
            # stay aligned regardless of which column carries an image.
            signature_table = next(
                t for t in document.tables
                if any("NGƯỜI YÊU CẦU" in cell.text for cell in t.rows[0].cells)
            )
            self.assertEqual(len(signature_table.rows), 2)
            headers = [cell.text.splitlines()[0] for cell in signature_table.rows[0].cells]
            self.assertEqual(headers, ["NGƯỜI YÊU CẦU", "CHỦ THUÊ BAO MỚI", "GIAO DỊCH VIÊN"])
            names = [cell.text for cell in signature_table.rows[1].cells]
            # Printed names in this table sit in the "Great Vibes" cursive
            # signature font, which title-cases them at substitution time
            # (all-caps collides in a script face) -- see renderer.py's
            # `_replace_placeholders` SIGNATURE_FONT_NAME branch.
            self.assertIn("Nguyễn Văn An", names[0])  # requester = individual old owner
            self.assertIn("Trần Thị B", names[1])  # new subscriber owner
            # The final section (previously forced to 3 columns purely to
            # fake this layout) is a normal single column now.
            final_columns = document._element.body.sectPr.xpath("./w:cols")
            if final_columns:
                self.assertEqual(
                    final_columns[0].get(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num"
                    ),
                    "1",
                )

    def test_aftersale_template_keeps_equal_marks_and_readable_body_spacing(self):
        data = self.report(DocumentType.AFTERSALE)
        document = Document(str(self.registry.for_data(data).template))
        mark_names = {
            "id_attachment_mark", "update_information_mark",
            "replace_sim_mark", "transfer_mark",
        }
        mark_runs = {
            name: run
            for paragraph in document.paragraphs
            for run in paragraph.runs
            for name in mark_names
            if name in run.text
        }
        self.assertEqual(set(mark_runs), mark_names)
        self.assertTrue(all(run.font.size and run.font.size.pt == 12 for run in mark_runs.values()))

        long_paragraphs = [
            paragraph for paragraph in document.paragraphs
            if paragraph.text.startswith((
                "Tôi xin cam kết là chủ sở hữu",
                "Tôi đồng ý thanh lý Hợp đồng",
                "Tôi ({{ requester_role_mark }}",
                "Tôi xin cam đoan số thuê bao",
                "Tôi đồng ý để Vietnamobile thu hồi",
            ))
        ]
        self.assertEqual(len(long_paragraphs), 6)
        for paragraph in long_paragraphs:
            spacing = paragraph._p.pPr.find(qn("w:spacing"))
            self.assertEqual(spacing.get(qn("w:line")), "240")
            self.assertEqual(spacing.get(qn("w:lineRule")), "auto")
            self.assertEqual(paragraph._p.pPr.find(qn("w:jc")).get(qn("w:val")), "both")

    def test_prepaid_organization_fills_only_organization_section(self):
        with tempfile.TemporaryDirectory() as folder:
            data = self.report(DocumentType.PREPAID_CONTRACT)
            data.customer = PersonData(
                entity_type="Tổ chức",
                organization_name="CÔNG TY TNHH VIỄN THÔNG MẪU",
                headquarters_address="10 Xã Đàn, Hà Nội",
                business_registration_number="0101234567",
                business_registration_issue_place="Sở Tài chính Hà Nội",
                business_registration_issue_date="01/01/2020",
                representative_name="NGUYỄN ĐẠI DIỆN",
                representative_position="Giám đốc",
                id_number="001090001234",
                date_of_birth="01/01/1990",
                issue_date="01/01/2022",
                issue_place="Cục Cảnh sát QLHC về TTXH",
                phone="0901234567",
            )
            output = self.registry.for_data(data).generate(data, Path(folder))
            text = docx_text(output)
            self.assertIn("CÔNG TY TNHH VIỄN THÔNG MẪU", text)
            self.assertIn("NGUYỄN ĐẠI DIỆN", text)

    def test_structured_prepaid_renders_company_representative_customer_and_sim_rows(self):
        with tempfile.TemporaryDirectory() as folder:
            data = self.report(DocumentType.PREPAID_CONTRACT)
            data.prepaid_structured_parties = True
            data.customer = PersonData(
                entity_type="Tổ chức", organization_name="CÔNG TY KHÁCH HÀNG",
                headquarters_address="10 Trần Phú", business_registration_number="0101234567",
                business_registration_issue_place="Sở KHĐT Hà Nội",
                business_registration_issue_date="01/01/2020",
            )
            data.representative = PersonData(
                full_name="LÊ NGƯỜI ĐẠI DIỆN", id_number="001090001234",
                date_of_birth="01/01/1990", issue_date="01/01/2022",
                issue_place="Cục Cảnh sát QLHC về TTXH", nationality="Việt Nam",
                address="10 Trần Phú", phone="0909888777", email="rep@example.com",
                representative_position="Giám đốc",
            )
            data.new_owner = PersonData(
                full_name="TRẦN KHÁCH HÀNG", id_number="001204001289",
                date_of_birth="01/02/2004", issue_date="02/02/2022",
                issue_place="Cục Cảnh sát QLHC về TTXH", nationality="Việt Nam",
                address="20 Hoàn Kiếm", phone="0987654321", email="customer@example.com",
            )
            data.provider_unit_address = "Địa chỉ đơn vị cung cấp"
            data.provider_representative = "ĐẠI DIỆN NHÀ CUNG CẤP"
            data.service_point_address = "Địa điểm giao dịch riêng"
            data.service_point_phone = "02435730123"
            data.prepaid_subscribers = [
                {"subscriber_number": "0925123456", "sim_serial": "8984041111111111111", "activation_date": "11/08/2026"},
                {"subscriber_number": "0987654321", "sim_serial": "8984042222222222222", "activation_date": "12/08/2026"},
                {"subscriber_number": "0909123456", "sim_serial": "8984043333333333333", "activation_date": "13/08/2026"},
            ]

            errors = self.registry.for_data(data).check(data)
            self.assertEqual(errors, [])
            output = self.registry.for_data(data).generate(data, Path(folder))
            text = docx_text(output)
            self.assertIn("CÔNG TY KHÁCH HÀNG", text)
            self.assertIn("LÊ NGƯỜI ĐẠI DIỆN", text)
            self.assertIn("TRẦN KHÁCH HÀNG", text)
            self.assertIn("Địa chỉ đơn vị cung cấp", text)
            self.assertIn("Địa điểm giao dịch riêng", text)
            self.assertIn("8984041111111111111", text)
            self.assertIn("8984042222222222222", text)
            self.assertIn("8984043333333333333", text)
            self.assertEqual(text.count("001090001234"), 1)
            self.assertNotIn("NGUYỄN VĂN AN", text)
            self.assertNotIn("{{", text)

    def test_prepaid_shop_phone_is_dynamic_and_provider_contacts_stay_internal(self):
        with tempfile.TemporaryDirectory() as folder:
            data = self.report(DocumentType.PREPAID_CONTRACT)
            data.provider_phone = "0909888777"
            data.provider_email = "daidien@example.com"
            data.shop_phone_2 = "0987000111"
            output = self.registry.for_data(data).generate(data, Path(folder))
            text = docx_text(output)
            self.assertNotIn("0909888777", text)
            self.assertNotIn("daidien@example.com", text)
            # The newly saved template keeps the provider hotline fixed and
            # uses the dynamic phone placeholder for the service point.
            self.assertIn("Điện thoại: (024) 35730123", text)
            self.assertIn(
                "Số điện thoại của điểm giao dịch: 02435730123 - 0987000111",
                text,
            )

    def test_prepaid_foreign_person_prints_passport_country(self):
        with tempfile.TemporaryDirectory() as folder:
            data = self.report(DocumentType.PREPAID_CONTRACT)
            data.customer.nationality = "Nhật Bản"
            data.customer.foreign_country = "Nhật Bản"
            output = self.registry.for_data(data).generate(data, Path(folder))
            text = docx_text(output)
            self.assertIn("☐ Việt Nam    ☑ Nước ngoài: Nhật Bản", text)

    def test_placeholder_split_across_word_runs_keeps_other_run_styles(self):
        from desktop_app.backend.documents.renderer import _replace_placeholders

        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("Họ tên: ").bold = True
        first = paragraph.add_run("{{ customer_")
        first.italic = True
        last = paragraph.add_run("name }} / giữ nguyên")
        last.underline = True

        _replace_placeholders(paragraph, {"customer_name": "NGUYỄN VĂN AN"})

        # _replace_placeholders only substitutes text now -- bold/size/dots
        # are baked directly into each template's own placeholder runs
        # instead (scripts/bake_field_highlighting.py), so a hand-built
        # paragraph like this one (no baked formatting of its own) gets a
        # plain substitution, keeping whatever formatting was already on
        # each surrounding run untouched.
        self.assertEqual(paragraph.text, "Họ tên: NGUYỄN VĂN AN / giữ nguyên")
        self.assertTrue(paragraph.runs[0].bold)
        self.assertTrue(paragraph.runs[1].italic)


if __name__ == "__main__":
    unittest.main()
